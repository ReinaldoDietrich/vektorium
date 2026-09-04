# -*- coding: utf-8 -*-
"""Tabelas da proposta em TABELA NATIVA do Word, reproduzindo as tabelas das telas. Fonte Arial
Narrow 8, tabela centralizada e AutoFit to Contents via XML OOXML (tblLayout autofit + tcW
type=auto w=0 em cada célula). Ref: ECMA-376 §17.4.53 (tblLayout), §17.4.72 (tcW)."""
from docx.shared import Pt, RGBColor, Twips
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from ..exportacao.compilacao_export import (_agrupar, COLUNAS_BASE, ELET_LABELS,
                                            COLUNAS_RESUMO_DEFAULT, COLUNAS_RESUMO_COMPRESSAO_DEFAULT)
from ..exportacao.consumo_export import COLUNAS as CONSUMO_COLUNAS, _linha_cenario
from ..exportacao.compilacao_geral_export import (BLOCO_CARGA, BLOCO_DADOS_ENTRADA,
                                                  BLOCO_FORCADORES, BLOCO_RACK, BLOCO_CONDENSADOR)

FONTE = 8
FONTE_NOME = "Arial Narrow"
COR_DARK = "111827"
COR_SUB = "374151"
COR_LIGHT = "D9D9D9"
COR_ZEBRA = "F3F4F6"
COR_PRETO = "000000"

# Estudo Luminotécnico — 11 colunas fixas em retrato (decisão do usuário 2026-08-18).
LUMINO_COLUNAS = [
    ("linha_succao", "Linha de Sucção"), ("ambiente", "Ambiente"), ("area", "Área (m²)"),
    ("plano_calculo", "Plano de Cálculo (m)"), ("qtd_luminarias", "Qtd Luminárias"),
    ("potencia_w", "Potência (W)"), ("fluxo_total", "Fluxo Total"),
    ("eficiencia_lmxw", "Eficiência Luminosa (LmxW)"), ("potencia_total_w", "Potência Total (W)"),
    ("densidade_w_m2", "Densidade (Wxm²)"), ("lux_calculado", "Lux Calculado"),
]



def _fmt(n, dec=0):
    if n is None or n == "":
        return "—"
    try:
        return f"{float(n):,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return str(n)


def _txt(v):
    return "—" if v is None or v == "" else str(v)


def _shd(cell, cor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), cor)
    tcPr.append(shd)


def _set(cell, texto, *, bold=False, white=False, align="c", size=FONTE):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                   "r": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    r = p.add_run("" if texto is None else str(texto))
    r.bold = bold
    r.font.name = FONTE_NOME
    r.font.size = Pt(size)
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONTE_NOME)
    if white:
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _tabela(sd, ncols):
    t = sd.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    return t


# ---------------------------------------------------------------- AutoFit to Contents (OOXML)
def aplicar_autofit(t, conteudo_por_coluna=None):
    """AutoFit to Contents via XML OOXML: tblLayout type=autofit + tcW type=auto w=0 em cada célula.
    Ref: ECMA-376 §17.4.53 (tblLayout), §17.4.72 (tcW), officeopenxml.com, python-docx issue #209.
    O parâmetro conteudo_por_coluna é mantido na assinatura por compatibilidade mas não é usado."""
    t.autofit = True                # tblLayout type="autofit" (via python-docx)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in t.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
            tcW.set(qn("w:w"), "0")
            tcW.set(qn("w:type"), "auto")
    _manter_tabela_junta(t)


def _manter_tabela_junta(t):
    """Não cortar linha entre páginas (cantSplit, ECMA-376 §17.4.7), repetir o cabeçalho (tblHeader),
    manter linhas juntas via keepNext (§17.3.1.15) e keepLines (§17.3.1.14) em TODAS as linhas."""
    nrows = len(t.rows)
    for i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))
        if i == 0 and trPr.find(qn("w:tblHeader")) is None:
            trPr.append(OxmlElement("w:tblHeader"))
        for par in row.cells[0].paragraphs:
            pPr = par._p.get_or_add_pPr()
            if pPr.find(qn("w:keepLines")) is None:
                pPr.append(OxmlElement("w:keepLines"))
            if i < nrows - 1 and pPr.find(qn("w:keepNext")) is None:
                pPr.append(OxmlElement("w:keepNext"))


def _novo_conteudo(rotulos):
    return [[r] for r in rotulos]


def _ac(conteudo, valores):
    for i, v in enumerate(valores):
        if i < len(conteudo):
            conteudo[i].append(v)


def _cabecalho(t, rotulos, esq_ate=1):
    row = t.add_row().cells
    for i, h in enumerate(rotulos):
        _set(row[i], h, bold=True, white=True, align="l" if i < esq_ate else "c")
        _shd(row[i], COR_DARK)


def _barra(t, texto, ncols, cor, white=True):
    row = t.add_row()
    c = row.cells[0]
    for i in range(1, ncols):
        c = c.merge(row.cells[i])
    _set(c, texto, bold=True, white=white, align="l")
    _shd(c, cor)


def _titulo(sd, texto):
    p = sd.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:keepNext"))     # título colado na tabela seguinte
    r = p.add_run(texto)
    r.bold = True
    r.font.name = FONTE_NOME
    r.font.size = Pt(11)
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONTE_NOME)


_SEQ_TABELA = [0]

def _titulo_abnt(sd, texto):
    _SEQ_TABELA[0] += 1
    p = sd.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:keepNext"))
    pPr.append(OxmlElement("w:keepLines"))
    r = p.add_run(f"Tabela {_SEQ_TABELA[0]} — {texto}")
    r.bold = True
    r.font.name = FONTE_NOME
    r.font.size = Pt(10)
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONTE_NOME)



def _cell_val(v):
    if v is None or v == "":
        return "—"
    if isinstance(v, float):
        return _fmt(v, 0) if float(v).is_integer() else _fmt(v, 1)
    return str(v)


def _linha_dados(t, valores, alinhas):
    row = t.add_row().cells
    for i, v in enumerate(valores):
        _set(row[i], v, bold=(i == 0), align=alinhas[i])


# ----------------------------------------------------------------- Compilação de Linhas (Tela 5)
def comp_linhas(sd, cl):
    itens = cl.get("itens", []) or []
    grupos, _gt, _ge = _agrupar(itens)

    _titulo_abnt(sd,"Distribuição de Linhas — Dados Mecânicos")
    nc = len(COLUNAS_BASE)
    t = _tabela(sd, nc)
    _cabecalho(t, list(COLUNAS_BASE), esq_ate=2)
    alin = ["l", "l", "c", "c", "l", "c", "c", "l"]
    cont = _novo_conteudo(list(COLUNAS_BASE))
    for g in grupos:
        _barra(t, f"SISTEMA {g['sistema']} — Temp. Evaporação: {_fmt(g['temp_evaporacao'],1)}°C — Gás: {g['gas_refrigerante']}", nc, COR_DARK)
        for r in g["ramais"]:
            _barra(t, f"Linha de Sucção {r['ramal']}", nc, COR_LIGHT, white=False)
            for c in r["itens"]:
                vals = [c["codigo"], c["nome"], c["modulacao_area"], f"{_fmt(c.get('carga_termica'))} kcal/h",
                        c["forcadores"], c["folga_real"], c["trocas_ar"], c["valvulas"]]
                _linha_dados(t, vals, alin); _ac(cont, vals)
            _barra(t, f"Subtotal Linha {r['ramal']}: {_fmt(r['subtotal_carga'])} kcal/h · {_fmt(r['subtotal_eletrica'])} W", nc, COR_ZEBRA, white=False)
        _barra(t, f"TOTAL SISTEMA {g['sistema']}: {_fmt(g['total_carga'])} kcal/h", nc, COR_DARK)
    aplicar_autofit(t, cont)

    sd.add_paragraph()
    _titulo_abnt(sd,"Distribuição de Linhas — Dados Elétricos")
    cols_e = ["Sistema", "Gabinete/Câmara"]
    for lbl, _, _, _ in ELET_LABELS:
        cols_e += [f"{lbl} (W)"]
    nce = len(cols_e)
    te = _tabela(sd, nce)
    _cabecalho(te, cols_e, esq_ate=2)
    aline = ["l", "l"] + ["c"] * (nce - 2)
    conte = _novo_conteudo(cols_e)
    resumo_por_sistema = {r["sistema_nome"]: r for r in (cl.get("resumo_sistemas") or [])}
    for g in grupos:
        _barra(te, f"SISTEMA {g['sistema']} — Temp. Evaporação: {_fmt(g['temp_evaporacao'],1)}°C — Gás: {g['gas_refrigerante']}", nce, COR_DARK)
        for r in g["ramais"]:
            _barra(te, f"Linha de Sucção {r['ramal']}", nce, COR_LIGHT, white=False)
            for c in r["itens"]:
                vals = [c["codigo"], c["nome"]]
                for _, w, a, dj in ELET_LABELS:
                    vals += [_fmt(c.get(w))]
                _linha_dados(te, vals, aline); _ac(conte, vals)
        rs = resumo_por_sistema.get(g["sistema"], {})
        _barra(te, f"TOTAL SISTEMA {g['sistema']}: Potência Máxima Instalada {_fmt(rs.get('potencia_total_w'))} W · Potência em Operação {_fmt(rs.get('potencia_demandada_w'))} W", nce, COR_DARK)
    aplicar_autofit(te, conte)

    _resumo_potencia(sd, "Resumo de Potência por Sistema — Quadros de Linhas",
                     cl.get("colunas_resumo") or [{"rotulo": r[1], "campo": r[2]} for r in COLUNAS_RESUMO_DEFAULT],
                     cl.get("resumo_sistemas") or [], "Quadro de Linhas - {sistema_nome}", "Sistema")
    _resumo_potencia(sd, "Resumo de Potência por Sistema — Compressão e Condensação",
                     cl.get("colunas_resumo_compressao") or [{"rotulo": r[1], "campo": r[2]} for r in COLUNAS_RESUMO_COMPRESSAO_DEFAULT],
                     cl.get("resumo_compressao") or [], "{sistema_nome} — {equipamento}", "Sistema / Equipamento")


def _resumo_potencia(sd, titulo, colunas, linhas, rotulo_fmt, rotulo_col1):
    if not linhas:
        return
    sd.add_paragraph()
    _titulo_abnt(sd,titulo)
    cols = [rotulo_col1] + [c["rotulo"] for c in colunas]
    nc = len(cols)
    t = _tabela(sd, nc)
    _cabecalho(t, cols, esq_ate=1)
    cont = _novo_conteudo(cols)
    for r in linhas:
        row = t.add_row().cells
        c1 = rotulo_fmt.format(**r)
        _set(row[0], c1, align="l")
        vals = [c1]
        for i, c in enumerate(colunas, start=1):
            v = _fmt(r.get(c["campo"]), 2) if str(c["campo"]).endswith("kva") else _fmt(r.get(c["campo"]))
            _set(row[i], v, align="c"); vals.append(v)
        _ac(cont, vals)
    aplicar_autofit(t, cont)


# ----------------------------------------------------------------- Compilação Geral (4 tabelas)
_DADOS_ENTRADA_PROPOSTA = [(c, l) for c, l in BLOCO_DADOS_ENTRADA if c not in ("equipamentos", "portas")]


def _cab_camaras(sd, camaras):
    nc = 1 + len(camaras)
    t = _tabela(sd, nc)
    row = t.add_row().cells
    _set(row[0], "", white=True); _shd(row[0], COR_DARK)
    for i, c in enumerate(camaras, start=1):
        _set(row[i], c["nome"], bold=True, white=True, align="c"); _shd(row[i], COR_DARK)
    return t, nc


def comp_geral(sd, dados):
    sistemas = [s for s in (dados.get("sistemas") or []) if s.get("camaras")]
    for s in sistemas:
        camaras = s["camaras"]
        cab = [""] + [c["nome"] for c in camaras]

        def linha_cam(t, cont, label, bloco_key, campo):
            r = t.add_row().cells
            _set(r[0], label, bold=True, align="l")
            vals = [label]
            for i, c in enumerate(camaras, start=1):
                v = _cell_val((c.get(bloco_key) or {}).get(campo))
                _set(r[i], v, align="c"); vals.append(v)
            _ac(cont, vals)

        def linha_sis(t, nc, cont, label, bloco, campo):
            r = t.add_row().cells
            _set(r[0], label, bold=True, align="l")
            v = _cell_val((bloco or {}).get(campo))
            for i in range(1, nc):
                _set(r[i], v, align="c")
            _ac(cont, [label] + [v] * (nc - 1))

        _titulo_abnt(sd,f"Sistema {s['sistema_nome']} — Informações de Carga e Dados de Entrada")
        t1, nc1 = _cab_camaras(sd, camaras); c1 = _novo_conteudo(cab)
        _barra(t1, "Informações Carga", nc1, COR_SUB)
        for campo, label in BLOCO_CARGA:
            linha_cam(t1, c1, label, "bloco_carga", campo)
        _barra(t1, "Dados de Entrada — Produto e Operação (fornecidos pelo cliente)", nc1, COR_SUB)
        for campo, label in _DADOS_ENTRADA_PROPOSTA:
            linha_cam(t1, c1, label, "dados_entrada", campo)
        aplicar_autofit(t1, c1); sd.add_paragraph()

        _titulo_abnt(sd,f"Sistema {s['sistema_nome']} — Informações Forçador de Ar")
        t2, nc2 = _cab_camaras(sd, camaras); c2 = _novo_conteudo(cab)
        for campo, label in BLOCO_FORCADORES:
            linha_cam(t2, c2, label, "forcador", campo)
        aplicar_autofit(t2, c2); sd.add_paragraph()

        _titulo_abnt(sd,f"Sistema {s['sistema_nome']} — Informações Unidades/Rack")
        t3, nc3 = _cab_camaras(sd, camaras); c3 = _novo_conteudo(cab)
        for campo, label in BLOCO_RACK:
            linha_sis(t3, nc3, c3, label, s.get("rack_uc"), campo)
        aplicar_autofit(t3, c3); sd.add_paragraph()

        _titulo_abnt(sd,f"Sistema {s['sistema_nome']} — Informações Condensadores")
        t4, nc4 = _cab_camaras(sd, camaras); c4 = _novo_conteudo(cab)
        for campo, label in BLOCO_CONDENSADOR:
            linha_sis(t4, nc4, c4, label, s.get("condensador"), campo)
        aplicar_autofit(t4, c4); sd.add_paragraph()


# ----------------------------------------------------------------- Consumo (Tela 9)
def consumo(sd, dados):
    cols = [r for r in CONSUMO_COLUNAS]
    for s in dados.get("sistemas", []) or []:
        if not s.get("disponivel", True):
            continue
        _titulo_abnt(sd,f"{s['sistema_nome']} — {_txt(s.get('modelo_uc'))}")
        t = _tabela(sd, len(cols)); cont = _novo_conteudo(cols)
        _cabecalho(t, cols, esq_ate=1)
        for nome_cen, chave in (("Projeto", "projeto"), ("Simples", "simples")):
            vals = _linha_cenario(nome_cen, s.get(chave) or {})
            row = t.add_row().cells
            saida = [vals[0]]
            _set(row[0], vals[0], align="l")
            for i, v in enumerate(vals[1:], start=1):
                txt = _fmt(v, 0) if i < 7 else ("R$ " + _fmt(v, 2))
                _set(row[i], txt, align="c"); saida.append(txt)
            _ac(cont, saida)
        aplicar_autofit(t, cont)
    tot = dados.get("total") or {}
    if tot:
        _titulo_abnt(sd,"Total do Projeto")
        cols2 = ["Cenário", "Total (kWh/mês)", "Custo/mês (R$)"]
        t = _tabela(sd, 3); cont = _novo_conteudo(cols2)
        _cabecalho(t, cols2, esq_ate=1)
        for nome, kk, ck in (("Projeto", "kwh_projeto_mes", "custo_projeto_mes"), ("Simples", "kwh_simples_mes", "custo_simples_mes")):
            row = t.add_row().cells
            a, b, c = nome, _fmt(tot.get(kk)), "R$ " + _fmt(tot.get(ck), 2)
            _set(row[0], a, align="l"); _set(row[1], b, align="c"); _set(row[2], c, align="c")
            _ac(cont, [a, b, c])
        aplicar_autofit(t, cont)


# ----------------------------------------------------------------- Luminotécnico (Tela 11)
def lumino(sd, dados):
    cols = [rot for _, rot in LUMINO_COLUNAS]
    t = _tabela(sd, len(cols)); cont = _novo_conteudo(cols)
    _cabecalho(t, cols, esq_ate=2)
    for item in dados.get("linhas", []) or []:
        row = t.add_row().cells
        vals = []
        for i, (chave, _) in enumerate(LUMINO_COLUNAS):
            v = _cell_val(item.get(chave))
            _set(row[i], v, align="l" if chave in ("linha_succao", "ambiente") else "c"); vals.append(v)
        _ac(cont, vals)
    aplicar_autofit(t, cont)
    resumo = dados.get("resumo_por_modelo") or []
    if resumo:
        sd.add_paragraph()
        _titulo_abnt(sd,"Resumo Somatório — por Modelo de Luminária")
        cols2 = ["Modelo Luminária", "Fabricante", "Qtd. Total"]
        t2 = _tabela(sd, 3); cont2 = _novo_conteudo(cols2)
        _cabecalho(t2, cols2, esq_ate=2)
        for r in resumo:
            row = t2.add_row().cells
            a, b, c = r["modelo_luminaria"], _txt(r.get("fabricante")), _txt(r.get("qtd_total"))
            _set(row[0], a, align="l"); _set(row[1], b, align="l"); _set(row[2], c, align="c")
            _ac(cont2, [a, b, c])
        aplicar_autofit(t2, cont2)


# ----------------------------------------------------------------- Painéis e Portas (Tela 7)
def paineis(sd, resumo):
    def tab_itens(linhas, titulo):
        if not linhas:
            return
        _titulo_abnt(sd,titulo)
        cols = ["Item", "Unid.", "Quantidade"]
        t = _tabela(sd, 3); cont = _novo_conteudo(cols)
        _cabecalho(t, cols, esq_ate=1)
        for it in linhas:
            row = t.add_row().cells
            a, b, c = it["item"], _txt(it["unidade"]), _fmt(it["quantidade"], 2)
            _set(row[0], a, align="l"); _set(row[1], b, align="c"); _set(row[2], c, align="c")
            _ac(cont, [a, b, c])
        aplicar_autofit(t, cont)

    tab_itens(resumo.get("paineis_parede_teto") or [], "Painéis — Parede e Teto")
    tab_itens(resumo.get("isolamento_piso") or [], "Isolamento de Piso")
    portas = resumo.get("portas") or []
    if portas:
        _titulo_abnt(sd,"Portas")
        cols = ["Descrição", "Id.", "Unid.", "Quantidade", "Tensão", "Observação"]
        t = _tabela(sd, 6); cont = _novo_conteudo(cols)
        _cabecalho(t, cols, esq_ate=1)
        for p in portas:
            row = t.add_row().cells
            vals = [p["descricao"], _txt(p.get("id_porta")), _txt(p.get("unidade")),
                    _fmt(p.get("quantidade")), _txt(p.get("tensao")), _txt(p.get("observacoes"))]
            _set(row[0], vals[0], align="l"); _set(row[1], vals[1], align="c")
            _set(row[2], vals[2], align="c"); _set(row[3], vals[3], align="c")
            _set(row[4], vals[4], align="c"); _set(row[5], vals[5], align="l")
            _ac(cont, vals)
        aplicar_autofit(t, cont)


# ----------------------------------------------------------------- Orçamento (Tela 10, com filtro)
def orcamento(sd, orc, filtro=None):
    filtro = filtro or {}
    agrupar = filtro.get("agrupar") or "bloco_cc"
    exibir = filtro.get("exibir") or "individual"
    is_bloco_cc = agrupar == "bloco_cc"
    show_qtd = is_bloco_cc and not filtro.get("ocultar_qtd")

    cols = ["Descrição"] + (["Unid.", "Qtd."] if show_qtd else []) + ["Valor (R$)"]
    nc = len(cols); col_val = nc - 1
    t = _tabela(sd, nc); cont = _novo_conteudo(cols)
    _cabecalho(t, cols, esq_ate=1)

    def _linha(label, *, unid="", qtd="", valor="", bold=False):
        row = t.add_row().cells
        _set(row[0], label, bold=bold, align="l")
        vals = [label]
        if show_qtd:
            _set(row[1], unid, align="c"); _set(row[2], qtd, align="c"); vals += [unid, qtd]
        _set(row[col_val], valor, bold=bold, align="c"); vals.append(valor)
        _ac(cont, vals)

    total_geral = 0.0
    for bloco in orc.get("ordem", []) or []:
        db = orc["blocos"][bloco]
        itens = db.get("itens", []) or []
        extra = db.get("extra_por_cc", {}) or {}
        if not itens and not extra:
            continue
        total_geral += db.get("total", 0) or 0
        if is_bloco_cc:
            _barra(t, f"{bloco} — R$ {_fmt(db.get('total'), 2)}", nc, COR_DARK)
        else:
            _linha(bloco, valor="R$ " + _fmt(db.get("total"), 2), bold=True)
            continue
        grupos, ordem_cc = {}, []
        for it in itens:
            cc = it.get("cc") or "Sem Centro de Custo"
            if cc not in grupos:
                grupos[cc] = []; ordem_cc.append(cc)
            grupos[cc].append(it)
        for cc in extra:
            if cc not in grupos:
                grupos[cc] = []; ordem_cc.append(cc)
        for cc in ordem_cc:
            _barra(t, cc, nc, COR_LIGHT, white=False)
            for it in grupos[cc]:
                _linha(it["descricao"], unid=_txt(it.get("unidade")), qtd=_fmt(it.get("quantidade"), 2),
                       valor=("R$ " + _fmt(it.get("valor"), 2)) if exibir == "individual" else "")
            if exibir == "total_bloco":
                subtotal = sum(i["valor"] for i in grupos[cc]) + (extra.get(cc) or 0)
                _linha(f"Subtotal — {cc}", valor="R$ " + _fmt(subtotal, 2))
    _linha("Valor total da proposta", valor="R$ " + _fmt(orc.get("total_geral", total_geral), 2), bold=True)
    aplicar_autofit(t, cont)


# ----------------------------------------------------------------- Exclusões
def excecoes(sd, empresa, linhas):
    cols = ["Atividade", empresa or "Empresa contratada", "Cliente"]
    t = _tabela(sd, 3); cont = _novo_conteudo(cols)
    _cabecalho(t, cols, esq_ate=1)
    for r in linhas:
        row = t.add_row().cells
        a = r["atividade"]; b = "X" if r.get("empresa") else ""; c = "X" if r.get("cliente") else ""
        _set(row[0], a, align="l"); _set(row[1], b, align="c"); _set(row[2], c, align="c")
        _ac(cont, [a, b, c])
    aplicar_autofit(t, cont)


# ----------------------------------------------------------------- Faturamento
def faturamento(sd, pares):
    t = _tabela(sd, 2); cont = _novo_conteudo(["Campo", "Dado"])
    for rot, val in pares:
        row = t.add_row().cells
        _set(row[0], rot, bold=True, align="l"); _set(row[1], _txt(val), align="l")
        _ac(cont, [rot, _txt(val)])
    aplicar_autofit(t, cont)
