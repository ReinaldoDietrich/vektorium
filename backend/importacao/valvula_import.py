# -*- coding: utf-8 -*-
"""Importação de seleção de Válvulas de Expansão via planilha Excel (1 aba, "Valvulas") — fonte é
o relatório de dimensionamento (Danfoss/Fullgauge/Carel), transcrito por uma IA a partir do
documento de instruções (ver valvula_docx_template.py). Cada linha casa com uma opção de forçador
específica via Identificador (código da câmara + código curto do forçador, ex.: "HTA1AabF1")."""
import io
import re
import openpyxl
from sqlalchemy.orm import Session
from .. import models as m

ABA_VALVULAS = "Valvulas"

# Campos novos (Capacidade/Orificio/Tensao/Tipo_Motor) vêm do relatório de dimensionamento do
# fabricante; o que o relatório não trouxer é completado pela "Tela A - Tabelas de Válvulas de
# Expansão" (banco) — prioridade sempre do relatório. Controlador NÃO entra na importação:
# seleção manual em caixa de seleção (decisão do usuário).
COLUNAS_VALVULAS = [
    "Identificador", "Fabricante", "Tipo_Expansao", "Modelo_Selecao",
    "Capacidade_Unit_Kcal_h", "Orificio", "Carga_Abertura_Pct",
    "Conexao_Entrada", "Conexao_Saida", "Tensao", "Tipo_Motor",
]

_RE_IDENTIFICADOR = re.compile(r"^(.*?)(F\d+)$")


def gerar_template_valvulas() -> bytes:
    # Modelo_Selecao SEMPRE na nomenclatura do banco do aplicativo (Tela A - Tabelas de Válvulas
    # de Expansão): "SB130T", "SB520", "E2V 24", "TEN5"... — NUNCA o código do fornecedor.
    # Orificio (só termostática) = o NÚMERO do orifício da seleção (ex.: "TE 5 - 2" -> 2).
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ABA_VALVULAS
    ws.append(COLUNAS_VALVULAS)
    ws.append(["HTA1AabF1", "Fullgauge", "Eletrônica", "SB130T", 5200, None, 74.44,
               "3/8\"", "1/2\"", "12 Vdc <> 10%", "Unipolar"])
    ws.append(["MTC1AaF1", "Danfoss", "Termostática", "TEN2", 3100, "4", None,
               "3/8\"", "1/2\"", None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _completar_do_banco(db: Session, item: dict):
    """Completa campos mecânicos FALTANTES com a Tela A - Tabelas de Válvulas de Expansão —
    o relatório do fabricante tem prioridade (só preenche o que veio vazio; divergência mantém
    o relatório). Match por modelo contido no texto da seleção (relatórios trazem códigos como
    '04336 - SB130T')."""
    modelo_txt = (item.get("modelo_selecao") or "").upper()
    if not modelo_txt:
        return
    candidatos = db.query(m.TabelaValvulaExpansao).all()
    achado = None
    for v in sorted(candidatos, key=lambda v: -len(v.modelo)):  # match mais específico primeiro
        if v.modelo.upper() in modelo_txt:
            achado = v
            break
    if not achado:
        # Modelo do relatório não existe na Tela A - Tabelas de Válvulas de Expansão — o usuário
        # é avisado na prévia (e na seção da válvula) pra digitar manualmente ou trocar o modelo.
        item["modelo_nao_cadastrado"] = True
        return
    item["modelo_nao_cadastrado"] = False
    for campo_item, campo_banco in (("conexao_entrada", "conexao_entrada"), ("conexao_saida", "conexao_saida"),
                                     ("tensao", "tensao"), ("tipo_motor", "tipo_motor")):
        if not item.get(campo_item) and getattr(achado, campo_banco):
            item[campo_item] = getattr(achado, campo_banco)
            item.setdefault("completados_banco", []).append(campo_item)


def _ler_aba(wb, nome_aba, colunas):
    if nome_aba not in wb.sheetnames:
        return []
    ws = wb[nome_aba]
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return []
    cab = [str(c).strip() if c is not None else "" for c in linhas[0]]
    idx = {col: cab.index(col) for col in colunas if col in cab}
    out = []
    for linha in linhas[1:]:
        if not linha or all(v is None for v in linha):
            continue
        reg = {col: (linha[idx[col]] if col in idx and idx[col] < len(linha) else None) for col in colunas}
        out.append(reg)
    return out


def ler_planilha_valvulas(conteudo: bytes) -> list:
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
    return _ler_aba(wb, ABA_VALVULAS, COLUNAS_VALVULAS)


def _parse_identificador(identificador: str):
    """'HTA1AabF1' -> ('HTA1Aab', 'F1'). Sem match -> (None, None)."""
    if not identificador:
        return None, None
    m_ = _RE_IDENTIFICADOR.match(identificador.strip())
    if not m_:
        return None, None
    return m_.group(1), m_.group(2)


def _mapa_forcadores_do_projeto(db: Session, projeto_id: int) -> dict:
    """{'HTA1AabF1': ('completo', camara, forcador_row), ...} — cobre câmaras Completo e
    Simples do projeto inteiro num único mapa, já que o relatório pode trazer os dois tipos."""
    from ..routers.camaras_completo import _codigo as _codigo_completo
    from ..routers.camaras_simples import _codigo as _codigo_simples

    mapa = {}
    sistemas = db.query(m.SistemaRefrigeracao).filter_by(projeto_id=projeto_id).all()
    for s in sistemas:
        for c in s.camaras_completo:
            cod_camara = _codigo_completo(c)
            for f in c.forcadores:
                if f.codigo_curto:
                    mapa[cod_camara + f.codigo_curto] = ("completo", c, f)
        for c in s.camaras_simples:
            cod_camara = _codigo_simples(c)
            for f in c.forcadores:
                if f.codigo_curto:
                    mapa[cod_camara + f.codigo_curto] = ("simples", c, f)
    return mapa


def montar_preview_valvulas(db: Session, projeto_id: int, conteudo: bytes) -> dict:
    linhas = ler_planilha_valvulas(conteudo)
    if not linhas:
        return {"erro": "Aba 'Valvulas' não encontrada ou vazia. Use o template."}
    mapa = _mapa_forcadores_do_projeto(db, projeto_id)

    itens = []
    for linha in linhas:
        identificador = (linha.get("Identificador") or "").strip()
        achado = mapa.get(identificador)
        item = {
            "identificador": identificador,
            "fabricante": linha.get("Fabricante"),
            "tipo_expansao": linha.get("Tipo_Expansao"),
            "modelo_selecao": linha.get("Modelo_Selecao"),
            "capacidade_unit_kcal_h": linha.get("Capacidade_Unit_Kcal_h"),
            "orificio": linha.get("Orificio"),
            "carga_abertura_pct": linha.get("Carga_Abertura_Pct"),
            "conexao_entrada": linha.get("Conexao_Entrada"),
            "conexao_saida": linha.get("Conexao_Saida"),
            "tensao": linha.get("Tensao"),
            "tipo_motor": linha.get("Tipo_Motor"),
        }
        _completar_do_banco(db, item)
        if achado:
            tipo_camara, camara, forcador = achado
            # Reimportação: sinaliza se a válvula desse forçador JÁ tem seleção gravada — o
            # frontend pergunta ao usuário se quer sobrepor antes de confirmar.
            Modelo = m.ValvulaSelecaoCompleto if tipo_camara == "completo" else m.ValvulaSelecaoSimples
            existente = (db.query(Modelo).filter_by(forcador_selecao_id=forcador.id, considerado=True).first()
                         or db.query(Modelo).filter_by(forcador_selecao_id=forcador.id).first())
            item.update({
                "encontrado": True, "tipo_camara": tipo_camara, "camara_id": camara.id,
                "camara_nome": camara.nome, "forcador_id": forcador.id,
                "forcador_linha": forcador.linha.nome if forcador.linha else None,
                "forcador_fabricante": forcador.fabricante.nome if forcador.fabricante else None,
                "ja_preenchida": bool(existente and existente.modelo_selecao),
            })
        else:
            item.update({"encontrado": False, "tipo_camara": None, "camara_id": None, "forcador_id": None,
                          "ja_preenchida": False})
        itens.append(item)
    return {"itens": itens}


def salvar_valvulas(db: Session, itens: list, sobrepor: bool = True) -> dict:
    """sobrepor=False (escolha do usuário na reimportação): válvulas que JÁ têm seleção gravada
    são mantidas intactas ('mantidas'); só as vazias recebem os dados da planilha."""
    criadas = atualizadas = ignoradas = mantidas = 0
    for it in itens:
        if not it.get("forcador_id") or not it.get("tipo_camara"):
            ignoradas += 1
            continue
        Modelo = m.ValvulaSelecaoCompleto if it["tipo_camara"] == "completo" else m.ValvulaSelecaoSimples
        existente = (db.query(Modelo).filter_by(forcador_selecao_id=it["forcador_id"], considerado=True).first()
                     or db.query(Modelo).filter_by(forcador_selecao_id=it["forcador_id"]).first())
        if existente and existente.modelo_selecao and not sobrepor:
            mantidas += 1
            continue
        campos = dict(
            fabricante=it.get("fabricante") or "—", tipo_expansao=it.get("tipo_expansao") or "Eletrônica",
            modelo_selecao=it.get("modelo_selecao"), carga_abertura_pct=it.get("carga_abertura_pct"),
            conexao_entrada=it.get("conexao_entrada"), conexao_saida=it.get("conexao_saida"),
            capacidade_unit_kcal_h=it.get("capacidade_unit_kcal_h"), orificio=it.get("orificio"),
            tensao=it.get("tensao"), tipo_motor=it.get("tipo_motor"),
        )
        if existente:
            for k, v in campos.items():
                setattr(existente, k, v)
            atualizadas += 1
        else:
            db.add(Modelo(forcador_selecao_id=it["forcador_id"], considerado=True, **campos))
            criadas += 1
    db.commit()
    return {"ok": True, "criadas": criadas, "atualizadas": atualizadas, "ignoradas": ignoradas, "mantidas": mantidas}
