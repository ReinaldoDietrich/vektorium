"""Gera o "Documento de Importação — Unidade Condensadora Comercial" (.docx) — específico para a
Tela B. Mesmo princípio do documento dos forçadores: o usuário preenche o cabeçalho e cola os
prints das tabelas do catálogo, e leva numa conversa com o Claude para gerar o Excel de importação.
A diferença central: a capacidade da UC depende de DUAS temperaturas (ambiente × evaporação) e cada
célula tem um par (capacidade Q / potência consumida P).

Mantido como alternativa ao fluxo principal (catálogo completo em PDF + script de extração, ver
protocolo_extracao_uc.md) — útil quando só se tem prints/recortes do catálogo, não o PDF inteiro."""
import io
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

CABECALHO_CAMPOS = [
    "Fabricante da Unidade Condensadora", "Versão/Ano do Catálogo",
    "Coluna Dimensional = Comprimento", "Coluna Dimensional = Largura", "Coluna Dimensional = Altura",
]

SECOES_TABELAS = [
    ("1. Tabela de Capacidades (Mecânica)",
     "Cole o print da tabela de capacidade. Ela vem indexada por temperatura ambiente "
     "(32/35/38/43°C) e por temperatura de evaporação, com DOIS valores por célula: capacidade "
     "(Q, kcal/h) e potência consumida (P, kW). O rótulo do par muda por fabricante (Q/P, CC/PC, "
     "C.R./P.C.) — é sempre capacidade + potência."),
    ("2. Tabela Elétrica",
     "Cole o print da tabela elétrica. Do compressor use sempre MCC, RLA e LRA (Corrente Máxima "
     "de Operação, Corrente Nominal e Corrente de Rotor Bloqueado). Dos ventiladores: tensão, "
     "fases, frequência, corrente e quantidade."),
    ("3. Dados Físicos e Dimensionais",
     "Cole o print da tabela com conexões (líquido/sucção), tanque de líquido, nível de ruído, "
     "diâmetro dos ventiladores, dimensões (Comp/Larg/Alt — indique a letra da coluna no "
     "cabeçalho) e peso (líquido e bruto)."),
    ("4. Nomenclatura do Código de Compra",
     "Cole o print da tabela COMPLETA que decompõe o código de compra do fabricante (ex.: U S H "
     "MB 4 120 J T D 2 C C 2), com TODOS os campos — inclusive os que parecem opcionais/de "
     "escolha do comprador (fluxo de ar, linha de líquido, versão, opcionais). O significado de "
     "cada código é dado do fabricante e deve ser extraído por completo."),
    ("5. Informações Gerais (opcional, para proposta comercial)",
     "Escreva abaixo um texto curto de descrição comercial (o que aparecerá futuramente numa "
     "proposta ao cliente quando este catálogo for selecionado). A foto do equipamento NÃO vai "
     "por aqui — não existe caminho pra imagem chegar no aplicativo pela planilha; suba a foto "
     "direto na Tela B do catálogo já criado, depois da importação."),
]

INSTRUCOES_IA = """Instruções para quem for processar este documento (IA ou humano, sem precisar de contexto prévio):

Objetivo: ler o cabeçalho preenchido e as imagens/textos das seções, e gerar um arquivo .xlsx com esta estrutura (5 abas):

Aba "Unidades": Fabricante_UC, Nome_Catalogo, Modelo, Versao_Catalogo, Sistema, Gas, Tipo_Compressor, Fabricante_Compressor, Numero_Compressores, HP, Vent_Qtd, Conexao_Liquido, Conexao_Succao, Tanque_Liquido_L, Nivel_Ruido_dB, Ventilador_Diametro_mm, Comprimento_mm, Largura_mm, Altura_mm, Peso_Liquido_kg, Peso_Bruto_kg, Descricao_Comercial (uma linha por modelo × gás — dados FÍSICOS/DIMENSIONAIS, que não variam por tensão). Nome_Catalogo é OBRIGATÓRIO e identifica o documento de origem — a dupla Fabricante+Versão sozinha não diferencia catálogos publicados na mesma data (ex.: "US 10 a 66HP - 2 e 3 Compressores").
Aba "Eletricas": Modelo, Tensao, Fases, Frequencia, Modelo_Compressor, MCC_A, RLA_A, LRA_A, Vent_Tensao, Vent_Fases, Vent_Frequencia, Vent_Corrente_A (uma linha por Modelo × Tensão × Modelo de Compressor — o catálogo real pode ter mais de uma opção de compressor pro mesmo Modelo+Tensão, cada uma com corrente própria: gere UMA LINHA PARA CADA)
Aba "Capacidades": Modelo, Temp_Ambiente_C, Temp_Evaporacao_C, Capacidade_kcal_h, Potencia_kW (uma linha por combinação temp. ambiente × temp. evaporação de cada modelo)
Aba "Campos": Ordem, Nome_Campo, Modo, Campo_Busca_Sistema, Substitui_Coringa, Codigo_Fixo, Valor, Codigo — estrutura do código comercial (compra), motor genérico usado pelo app (ver campo_catalogo.py), um catálogo inteiro por planilha (sem prefixo Fabricante/Linha, diferente do Forçador). Uma linha POR OPÇÃO de cada campo; os metadados do campo (Ordem/Nome_Campo/Modo/Campo_Busca_Sistema/Substitui_Coringa/Codigo_Fixo) repetem em todas as linhas do mesmo Nome_Campo, só Valor/Codigo mudam por linha. Extraia TODOS os campos da tabela de nomenclatura (seção 4), mesmo os que representam escolha do comprador (Fluxo de Ar, Linha de Líquido, Versão, Opcionais) — o dicionário de código-para-significado é dado do fabricante, mesmo quando o VALOR só é escolhido pelo comprador na hora da compra. Regras de Modo:
  - "automatico": o valor vem de um dado já conhecido do app (Campo_Busca_Sistema: tensao_comando, tensao_equipamentos, gas, sistema, tipo_compressor, fabricante_compressor ou numero_compressores) — SEMPRE precisa de uma linha Valor->Codigo pra cada valor possível, mesmo quando o código É literalmente o próprio valor (não existe atalho "usa o valor bruto sem tabela").
  - "fixo": usa direto o Codigo_Fixo, sempre igual — não precisa de linha Valor/Codigo.
  - "manual": o usuário escolhe entre as opções (Valor->Codigo) na hora de montar o pedido de compra (Tela 1) — se não escolher, o código mostra "*" no lugar desse campo. É o Modo típico de Fluxo de Ar, Linha de Líquido, Versão, Opcionais (escolha do comprador).
  - "ignorar": campo aparece na tela mas não entra no código de compra — não precisa de linha Valor/Codigo.
  Substitui_Coringa (1/0): marque 1 só se esse campo deve substituir um caractere "*" já existente no próprio código do Modelo (aba Unidades) em vez de ser concatenado no final.
  CARD OBRIGATÓRIO "Modelo Pesquisa": a aba "Campos" deve incluir uma linha com Nome_Campo = "Modelo Pesquisa" e Modo = "modelo_pesquisa" (Campo_Busca_Sistema/Codigo_Fixo/Valor/Codigo em branco). É o card que representa o próprio código do modelo (coluna Modelo da aba "Unidades") dentro da nomenclatura de compra — o app usa esse valor automaticamente, não invente um Campo "prefixo fixo" pra simular isso. O Ordem é só ponto de partida, o usuário reordena depois na tela do catálogo. Nunca prefixe/altere o texto da coluna Modelo pra "simular" essa posição — transcreva exatamente como está no catálogo (com "*" coringa se o fabricante já usa, ex.: "U*HMB4250"; sem, se não usa).
  REGRA OBRIGATÓRIA DE MODO (pros demais campos, exceto "Modelo Pesquisa" acima): a IA NUNCA decide sozinha se um campo é "automatico", "fixo" ou "ignorar" — isso é decisão exclusiva do usuário, feita depois na tela do próprio catálogo do app (que já tem edição de nomenclatura). Todo campo da aba "Campos" deve ser inserido com Modo = "manual" (Campo_Busca_Sistema e Codigo_Fixo em branco), preenchendo só Valor/Codigo com o que estiver literalmente na tabela de nomenclatura do catálogo. Só use "automatico"/"fixo"/"ignorar" se o usuário pedir isso explicitamente nesta conversa — nunca por inferência própria.
Aba "Fatores_Gas": (opcional, se o catálogo trouxer fator de correção por gás) Fabricante_UC, Gas, Fator

Regra fixa de domínio — Tensão da Unidade Condensadora (compressor + ventilador): quando (e só quando) o usuário pedir explicitamente Modo "automatico" para o campo de Tensão/Código de tensão, o Campo_Busca_Sistema é SEMPRE "tensao_equipamentos" — Unidade Condensadora é equipamento de potência, nunca "tensao_comando". Regra fixa do domínio, não uma escolha da IA.

SEM AUTONOMIA: quem processa este documento (IA ou humano) não tem autonomia para tomar nenhuma decisão de mapeamento, valor ou estrutura além do que está literalmente escrito no cabeçalho e nas tabelas coladas pelo usuário, nem para alterar o fluxo ou a estrutura descrita nestas instruções. Qualquer ambiguidade, dado faltante ou decisão de engenharia (Modo, convenção de capacidade, mapeamento de coluna, etc.) deve ser perguntada ao usuário — nunca resolvida por conta própria, mesmo que pareça óbvia.

Regras importantes:
- A tabela de nomenclatura (seção 4) decompõe o código de compra em vários campos, cada um virando uma linha (ou várias, uma por opção) na aba "Campos". Não confundir com a aba "Unidades": o que NÃO se extrai por modelo é qual valor de cada campo se aplica a cada unidade específica (isso é escolha do usuário no app, Tela 1) — o dicionário completo de código-para-significado de cada campo é dado do fabricante e vai inteiro na aba "Campos".
- A tabela elétrica do catálogo real costuma trazer o MESMO Modelo+Tensão repetido para cada opção de compressor disponível (ex.: mesmo modelo/tensão com compressor H5O5CC OU H7O5CC, cada um com MCC/RLA/LRA diferentes) — NÃO resuma numa linha só: gere uma linha na aba "Eletricas" para CADA combinação Modelo × Tensão × Modelo_Compressor que aparecer.
- A capacidade da UC depende de DUAS temperaturas: ambiente (32/35/38/43) e evaporação. Cada célula do catálogo tem um par capacidade (Q) / potência consumida (P). Gere uma linha na aba "Capacidades" por combinação, preenchendo Capacidade_kcal_h (Q) e Potencia_kW (P).
- Do compressor, extraia SEMPRE MCC, RLA e LRA na aba "Eletricas". Se o catálogo só trouxer MCC: RLA = MCC/1,56 (para compressor Copeland, RLA = MCC/1,4). LRA vem direto do catálogo.
- Extraia também Numero_Compressores (quantidade de compressores da unidade — usado no cálculo de consumo elétrico da Tela D), na aba "Unidades". Não confundir com Vent_Qtd (quantidade de ventiladores).
- Físico/Dimensional (conexões, tanque, ruído, diâmetro de ventilador, dimensões, peso) NÃO varia por tensão nem por modelo de compressor no catálogo real — fica só na aba "Unidades", uma vez por modelo, mesmo que a aba "Eletricas" tenha várias linhas pra esse mesmo Modelo.
- As ÚNICAS colunas que podem faltar num catálogo e sobrar noutro são as temperaturas de evaporação (cada fabricante cobre seu envelope). Todas as demais colunas, se o catálogo não trouxer o dado, devem permanecer no cadastro porém EM BRANCO — nunca omita a coluna.
- Orientação #11: se um mesmo equipamento atende mais de um gás, DUPLIQUE o cadastro — uma linha por gás na aba "Unidades" (e as capacidades/elétricas correspondentes) — para evitar conflito. Isso vale mesmo que o código mecânico do catálogo seja idêntico entre os gases.
- Se dois fabricantes de compressor diferentes usarem o mesmo código de modelo mecânico (ex.: mesma linha atendida por compressor Bitzer OU Dorin), NÃO embuta a marca do compressor no campo Modelo (nunca "U*HMB4120 Bitzer") — o campo Fabricante_Compressor já resolve isso, e junto com Modelo+Gas forma a chave que evita conflito na importação.
- Nomes de modelo podem ter qualquer caractere (letras, números, *, #, etc.) — transcreva exatamente como aparece.
- Para a tabela dimensional: use os campos "Coluna Dimensional = Comprimento/Largura/Altura" do cabeçalho para mapear as letras genéricas (A, L, P...) da imagem. Se o cabeçalho tiver um NÚMERO em vez de uma letra (ex.: "Largura = 357"), isso significa que essa medida é FIXA para todos os modelos — use esse valor direto, sem tentar mapear letra nenhuma para essa medida.
- A seção 4 (Nomenclatura) e a seção 5 (Informações Gerais) são opcionais só no sentido de que o usuário pode deixá-las em branco (nesse caso gere a aba "Campos" vazia, só cabeçalho, e deixe "Descricao_Comercial" em branco em todas as linhas). Se o usuário escreveu o texto da seção 5, copie-o pra TODAS as linhas da aba "Unidades" (mesmo texto repetido, vale pro catálogo inteiro, não por modelo) — nunca invente texto se a seção estiver em branco.
- Sempre mostre os dados extraídos ao usuário em formato de tabela, para revisão e confirmação, antes de gerar o .xlsx final.
- CAMINHO OBRIGATÓRIO DE SALVAMENTO: salve o .xlsx em "B:\\Documentos Programas\\App Carga Térmica\\Documentos de Criação\\A - Catálogos\\2 - Unidades Condensadoras\\", nome "Importacao_UC_{Fabricante}_{Versao}.xlsx". Nunca em pasta temporária.
- O arquivo deve ser importado na tela "Tela B — Unidades Condensadoras" do aplicativo (botão de upload)."""


def _titulo(doc, texto, tamanho=16, cor=RGBColor(0x11, 0x18, 0x27)):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(tamanho)
    run.font.color.rgb = cor
    return p


def _placeholder(doc):
    tabela = doc.add_table(rows=1, cols=1)
    tabela.style = "Table Grid"
    celula = tabela.cell(0, 0)
    celula.text = "[Cole a imagem aqui]"
    for paragrafo in celula.paragraphs:
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragrafo.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
    tabela.rows[0].height = Inches(2.2)
    doc.add_paragraph()


def gerar_documento_importacao_uc() -> bytes:
    doc = Document()

    _titulo(doc, "Documento de Importação — Unidade Condensadora Comercial", 18)
    intro = doc.add_paragraph(
        "Preencha o cabeçalho abaixo e cole as imagens (prints) de cada tabela do catálogo nas "
        "seções indicadas. Depois, envie este arquivo numa conversa com o Claude pedindo para "
        "gerar a planilha Excel de importação — as instruções já estão neste documento."
    )
    intro.runs[0].font.size = Pt(11)
    doc.add_paragraph()

    _titulo(doc, "Cabeçalho", 13)
    tabela_cab = doc.add_table(rows=len(CABECALHO_CAMPOS), cols=2)
    tabela_cab.style = "Table Grid"
    for i, campo in enumerate(CABECALHO_CAMPOS):
        cel_campo = tabela_cab.cell(i, 0)
        cel_campo.text = campo
        cel_campo.paragraphs[0].runs[0].bold = True
        tabela_cab.cell(i, 1).text = ""
    doc.add_paragraph()

    for titulo_secao, instrucao in SECOES_TABELAS:
        _titulo(doc, titulo_secao, 13)
        doc.add_paragraph(instrucao)
        if titulo_secao.startswith("5."):
            # Seção 5 é só texto (descrição comercial) — sem placeholder de imagem, foto não vai
            # nessa planilha de jeito nenhum (mesmo padrão da seção 6 do Forçador).
            p = doc.add_paragraph()
            p.add_run("Descrição comercial: ").bold = True
            doc.add_paragraph("_" * 90)
        else:
            _placeholder(doc)

    doc.add_page_break()
    _titulo(doc, "Instruções para quem for processar este documento", 13, RGBColor(0xB4, 0x53, 0x09))
    for linha in INSTRUCOES_IA.split("\n"):
        p = doc.add_paragraph(linha or " ")
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
