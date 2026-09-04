"""Gera o "Documento de Importação" (.docx) — um arquivo autoexplicativo que o usuário preenche
(cabeçalho + cola prints das tabelas do catálogo) e leva para uma conversa com o Claude (esta ou
qualquer outra, sem precisar de contexto prévio) para gerar o Excel de importação. As instruções
embutidas no próprio arquivo são o que garante isso — não dependem de memória de conversa."""
import io
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

CABECALHO_CAMPOS = [
    "Fabricante", "Linha", "Versão/Ano do Catálogo", "Aletas por Polegada (FPI)",
    "DT de Catálogo (°C)", "PDL Máximo Recomendado (m)",
    "Coluna Dimensional = Comprimento", "Coluna Dimensional = Largura", "Coluna Dimensional = Altura",
]

# Cada fabricante publica seu próprio fator de correção de capacidade por gás refrigerante —
# preenchido manualmente pelo usuário (consultando o catálogo do fabricante), um valor por gás.
GASES_FATOR_CORRECAO = ["R-404A", "R-134a", "R-452A", "R-448A", "R-449A", "R-407C", "R-22", "R-410A"]

SECOES_TABELAS = [
    ("1. Tabela Mecânica (Capacidades)",
     "Cole abaixo o print da tabela de capacidades por temperatura de evaporação, vazão de ar, "
     "quantidade de ventiladores, diâmetro e flecha de ar."),
    ("2. Tabela Elétrica",
     "Cole abaixo o print da tabela de dados elétricos (degelo e motores, por tensão)."),
    ("3. Dados Físicos",
     "Cole abaixo o print da tabela de dados físicos (linha de líquido, linha de sucção, "
     "equalizador, dreno, peso líquido, carga de refrigerante)."),
    ("4. Dados Dimensionais",
     "Cole abaixo o print da tabela de dimensões. As colunas costumam vir com letras genéricas "
     "(A, B, C...) em vez de \"comprimento/largura/altura\" — preencha no cabeçalho acima qual "
     "letra corresponde a cada uma. Se alguma medida for a MESMA para todos os modelos (não muda "
     "linha a linha), digite o valor numérico direto no cabeçalho em vez de uma letra — isso "
     "sinaliza que é um valor fixo, não uma coluna da tabela. Inclua também o número de fixações, "
     "se houver."),
    ("5. Nomenclatura do Código de Compra",
     "Cole abaixo o print da tabela que decompõe o código de compra do fabricante em campos "
     "(ex.: modelo \"FLA-B-4-017\" onde B = tensão, 4 = aletas, etc.). Usada para montar o código "
     "comercial automaticamente na tela do sistema (Fixo/Automático/Manual, por campo) — sem "
     "cadastro manual depois. Se o catálogo não tiver essa tabela, deixe em branco."),
    ("6. Informações Gerais (opcional, para proposta comercial)",
     "Escreva abaixo um texto curto de descrição comercial (o que aparecerá futuramente numa "
     "proposta ao cliente quando este catálogo for selecionado). A foto do equipamento NÃO vai "
     "por aqui — não existe caminho pra imagem chegar no aplicativo pela planilha; suba a foto "
     "direto na Tela A/B do catálogo já criado, depois da importação."),
]

INSTRUCOES_IA = """Instruções para quem for processar este documento (IA ou humano, sem precisar de contexto prévio):

Objetivo: ler o cabeçalho preenchido (incluindo a tabela de Fator de Correção por Gás) e as
imagens/textos das 6 seções abaixo, e gerar um arquivo .xlsx com exatamente esta estrutura (7 abas):

Aba "Modelos": Fabricante, Linha, Versao_Catalogo, Modelo, FPI, Num_Ventiladores, Diametro_Ventilador_mm, Tipo_Degelo, Carga_Gas_kg, Pot_Resistencia_Degelo_W, Vazao_Ar_m3h, DT_Referencia_C, PDL_Referencia_m, Flecha_Ar_m, Coletores_Por_Forcador, Descricao_Comercial (uma linha por modelo)
Aba "Capacidades": Modelo, Temp_Evaporacao_C, Capacidade_kcal_h (uma linha por temperatura de cada modelo)
Aba "Eletrica": Modelo, Tensao, Degelo_W, Degelo_A, Motores_W, Motores_A (uma linha por tensão de cada modelo)
Aba "Fisicos": Modelo, Linha_Liquido, Linha_Succao, Equalizador, Dreno, Peso_Liquido_kg, Carga_Refrigerante_kg (uma linha por modelo)
Aba "Dimensionais": Modelo, Comprimento_mm, Largura_mm, Altura_mm, Num_Fixacoes (uma linha por modelo)
Aba "Campos": Fabricante, Linha, Ordem, Nome_Campo, Modo, Campo_Busca_Sistema, Substitui_Coringa, Codigo_Fixo, Valor, Codigo — estrutura do código comercial (compra), motor genérico usado pelo app (ver campo_catalogo.py). Uma linha POR OPÇÃO de cada campo; os metadados do campo (Ordem/Nome_Campo/Modo/Campo_Busca_Sistema/Substitui_Coringa/Codigo_Fixo) repetem em todas as linhas do mesmo Nome_Campo, só Valor/Codigo mudam por linha. Regras de Modo:
  - "automatico": o valor vem de um dado já conhecido do app (Campo_Busca_Sistema: tipo_degelo, tensao_comando, tensao_equipamentos, gas, num_ventiladores ou diametro_ventilador_mm) — SEMPRE precisa de uma linha Valor->Codigo pra cada valor possível, mesmo quando o código É literalmente o próprio valor (ex.: num_ventiladores: Valor "1" Codigo "1", Valor "2" Codigo "2"... não existe atalho "usa o valor bruto sem tabela").
  - "fixo": usa direto o Codigo_Fixo, sempre igual — não precisa de linha Valor/Codigo.
  - "manual": o usuário escolhe entre as opções (Valor->Codigo) na hora de montar o pedido de compra — se não escolher, o código mostra "*" no lugar desse campo.
  - "ignorar": campo aparece na tela mas não entra no código de compra — não precisa de linha Valor/Codigo.
  Substitui_Coringa (1/0): marque 1 só se esse campo deve substituir um caractere "*" já existente no próprio código do Modelo (aba Modelos) em vez de ser concatenado no final — só faz sentido se os modelos dessa linha tiverem "*" no nome (ex.: "FM*108").
  CARD OBRIGATÓRIO "Modelo Pesquisa": TODA linha da aba "Campos" deve incluir uma linha com Nome_Campo = "Modelo Pesquisa" e Modo = "modelo_pesquisa" (Campo_Busca_Sistema/Codigo_Fixo/Valor/Codigo em branco). É o card que representa o próprio número do modelo (coluna Modelo da aba "Modelos") dentro da nomenclatura de compra — o app usa esse valor automaticamente, não precisa (nem deve) inventar um Campo "prefixo fixo" pra tentar simular isso. O Ordem dele é só um ponto de partida — o usuário reordena livremente depois, na tela do catálogo. Nunca prefixe/altere o texto da coluna Modelo pra "simular" essa posição — transcreva o modelo exatamente como aparece no catálogo (com "*" coringa se o fabricante já usa um, ex.: "FM*108"; sem, se não usa, ex.: "0509").
  REGRA OBRIGATÓRIA DE MODO (pros demais campos, exceto "Modelo Pesquisa" acima): a IA NUNCA decide sozinha se um campo é "automatico", "fixo" ou "ignorar" — isso é decisão exclusiva do usuário, feita depois na tela do próprio catálogo do app (que já tem edição de nomenclatura). Todo campo da aba "Campos" deve ser inserido com Modo = "manual" (Campo_Busca_Sistema e Codigo_Fixo em branco), preenchendo só Valor/Codigo com o que estiver literalmente na tabela de nomenclatura do catálogo. Só use "automatico"/"fixo"/"ignorar" se o usuário pedir isso explicitamente nesta conversa — nunca por inferência própria.
Aba "Fatores_Gas": Fabricante, Linha, Gas, Fator (uma linha por gás da tabela "Fator de Correção por Gás Refrigerante" do cabeçalho; se o usuário deixou um gás em branco, não inclua a linha desse gás)

Regra fixa de domínio — Tensão do motor/degelo do forçador (evaporador): quando (e só quando) o usuário pedir explicitamente Modo "automatico" para o campo de Tensão, o Campo_Busca_Sistema é SEMPRE "tensao_comando" — evaporador segue a tensão de comando do projeto, nunca "tensao_equipamentos". Regra fixa do domínio, não uma escolha da IA.

SEM AUTONOMIA: quem processa este documento (IA ou humano) não tem autonomia para tomar nenhuma decisão de mapeamento, valor ou estrutura além do que está literalmente escrito no cabeçalho e nas tabelas coladas pelo usuário, nem para alterar o fluxo ou a estrutura descrita nestas instruções. Qualquer ambiguidade, dado faltante ou decisão de engenharia (Modo, convenção de capacidade, mapeamento de coluna, etc.) deve ser perguntada ao usuário — nunca resolvida por conta própria, mesmo que pareça óbvia.

Regras importantes:
- Os valores do Cabeçalho (Fabricante, Linha, Versão/Ano, FPI, DT de Catálogo, PDL Máximo) se aplicam a TODOS os modelos lidos — preencha-os identicamente em cada linha da aba "Modelos" (FPI -> FPI, DT de Catálogo -> DT_Referencia_C, PDL Máximo -> PDL_Referencia_m).
- "Peso" é sempre peso líquido (sem embalagem) — nunca peso bruto. Se a tabela trouxer os dois, use só o líquido.
- Nomes de modelo podem ter qualquer caractere (letras, números, *, #, etc.) — transcreva exatamente como aparece na imagem, não filtre por formato.
- Se a tabela mecânica trouxer mais de uma convenção de capacidade (ex.: "DTML" e "DT" simples), pergunte ao usuário qual usar — a preferência já estabelecida é a convenção "DT" simples, não "DTML", a menos que ele diga o contrário.
- Para a tabela dimensional: use os campos "Coluna Dimensional = Comprimento/Largura/Altura" do cabeçalho para saber qual letra genérica (A, B, C...) da imagem mapeia para qual medida. Se o usuário não preencheu esses campos e a tabela usa letras genéricas, pergunte antes de adivinhar. Se o cabeçalho tiver um NÚMERO em vez de uma letra (ex.: "Largura = 357"), isso significa que essa medida é FIXA para todos os modelos daquela linha — não vem da tabela, use esse valor direto na aba "Dimensionais" para todo modelo, em vez de tentar mapear letra nenhuma para essa medida.
- A seção 5 (Nomenclatura) e a seção 6 (Informações Gerais) são opcionais só no sentido de que o usuário pode deixá-las em branco (nesse caso gere a aba "Campos" vazia, só cabeçalho, e deixe "Descricao_Comercial" em branco). Mas se o usuário COLOU a imagem da tabela de nomenclatura, extraia TODOS os campos dela — inclusive os que parecem opcionais/de escolha do comprador (ex.: cor, acabamento, acessório, esses viram Modo "manual") — o dicionário de código-para-significado de cada campo é dado do fabricante, mesmo quando o VALOR de um campo específico só é escolhido pelo comprador na hora da compra. A foto do equipamento não vai no Excel de forma nenhuma — não anote local de imagem em campo nenhum, o Excel só carrega Descricao_Comercial (texto).
- Sempre mostre os dados extraídos ao usuário em formato de tabela, para revisão e confirmação, antes de gerar o arquivo .xlsx final — nunca gere o arquivo sem essa confirmação explícita.
- CAMINHO OBRIGATÓRIO DE SALVAMENTO: o arquivo .xlsx final deve ser salvo em "B:\\Documentos Programas\\App Carga Térmica\\Documentos de Criação\\A - Catálogos\\1 - Forçadores\\", nome do arquivo "Importacao_{Fabricante}_{Linha}.xlsx" — nunca em pasta temporária. Essa é a única pasta aceita, independentemente de qual conversa/sessão estiver processando este documento."""


def _adicionar_titulo(doc, texto, tamanho=16, cor=RGBColor(0x11, 0x18, 0x27)):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(tamanho)
    run.font.color.rgb = cor
    return p


def _adicionar_placeholder(doc):
    tabela = doc.add_table(rows=1, cols=1)
    tabela.style = "Table Grid"
    celula = tabela.cell(0, 0)
    celula.text = "[Cole a imagem aqui]"
    for paragrafo in celula.paragraphs:
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragrafo.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
    # força uma altura mínima generosa para caber a imagem colada
    tabela.rows[0].height = Inches(2.2)
    doc.add_paragraph()


def gerar_documento_importacao() -> bytes:
    doc = Document()

    _adicionar_titulo(doc, "Documento de Importação — Catálogo de Forçadores de Ar", 18)
    intro = doc.add_paragraph(
        "Preencha o cabeçalho abaixo e cole as imagens (prints) de cada tabela do catálogo nas "
        "seções indicadas. Depois, envie este arquivo numa conversa com o Claude pedindo para "
        "gerar a planilha Excel de importação do sistema Carga Térmica — não é preciso explicar "
        "nada além disso, as instruções já estão neste documento."
    )
    intro.runs[0].font.size = Pt(11)
    doc.add_paragraph()

    _adicionar_titulo(doc, "Cabeçalho", 13)
    tabela_cab = doc.add_table(rows=len(CABECALHO_CAMPOS), cols=2)
    tabela_cab.style = "Table Grid"
    for i, campo in enumerate(CABECALHO_CAMPOS):
        cel_campo = tabela_cab.cell(i, 0)
        cel_campo.text = campo
        cel_campo.paragraphs[0].runs[0].bold = True
        tabela_cab.cell(i, 1).text = ""
    doc.add_paragraph()

    _adicionar_titulo(doc, "Fator de Correção por Gás Refrigerante", 13)
    doc.add_paragraph(
        "Cada fabricante publica seu próprio fator de correção de capacidade por gás — consulte o "
        "catálogo do fabricante e preencha um valor por gás (deixe em branco o que não constar)."
    )
    tabela_gas = doc.add_table(rows=len(GASES_FATOR_CORRECAO), cols=2)
    tabela_gas.style = "Table Grid"
    for i, gas in enumerate(GASES_FATOR_CORRECAO):
        cel_gas = tabela_gas.cell(i, 0)
        cel_gas.text = gas
        cel_gas.paragraphs[0].runs[0].bold = True
        tabela_gas.cell(i, 1).text = ""
    doc.add_paragraph()

    for titulo_secao, instrucao in SECOES_TABELAS:
        _adicionar_titulo(doc, titulo_secao, 13)
        doc.add_paragraph(instrucao)
        if titulo_secao.startswith("6."):
            # Seção 6 é só texto (descrição comercial) — sem placeholder de imagem, foto não vai
            # nessa planilha de jeito nenhum.
            p = doc.add_paragraph()
            p.add_run("Descrição comercial: ").bold = True
            doc.add_paragraph("_" * 90)
        else:
            _adicionar_placeholder(doc)

    doc.add_page_break()
    _adicionar_titulo(doc, "Instruções para quem for processar este documento", 13, RGBColor(0xB4, 0x53, 0x09))
    for linha in INSTRUCOES_IA.split("\n"):
        p = doc.add_paragraph(linha or " ")
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
