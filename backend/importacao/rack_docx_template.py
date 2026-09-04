# -*- coding: utf-8 -*-
"""Gera os "Documentos de Importação — Rack Paralelo" (.docx), um por fabricante de compressor
(Bitzer e Copeland têm formatos de relatório de seleção bem diferentes). O usuário anexa o PDF do
relatório junto com o documento correspondente numa conversa com o Claude (ou outra IA), que
devolve a planilha .xlsx pronta para upload na Tela 6 — mesmo padrão de Válvulas de Expansão."""
import io
from docx import Document
from docx.shared import Pt, RGBColor

COLUNAS = ("Sistema, Quantidade_Compressores, Tensao, Linha_Compressor, "
           "Fabricante_Compressor, Modelo_Compressor, COP_Compressor, Capacidade_Compressor_Kcal_H, "
           "Carga_Total_Fornecida_Kcal_H, Calor_Total_Rejeitado_Kcal_H, Potencia_Total_W, "
           "Corrente_Nominal_A, Corrente_Maxima_Trabalho_A, Vazao_Massica_Kg_H, Carga_Oleo_L, "
           "Conexao_Descarga, Conexao_Succao, HP_Total, Tanque_Liquido_L, Carga_Gas_Estimada_Kg")

REGRAS_COMUNS = """Regras importantes (valem para os dois fabricantes):
- NUNCA invente dado que não está no relatório. Campo que o relatório não traz: preencha com o texto "Não informado" — nunca deixe a célula vazia nem chute um valor.
- Sistema: o nome EXATO do sistema de refrigeração do projeto no aplicativo (ex.: "HTA", "MTB"), como o usuário informar na conversa — não é um dado do PDF do fabricante. Pergunte ao usuário qual sistema esse relatório representa se não estiver claro pelo contexto da conversa.
- Um PDF pode conter mais de uma seleção (rack + acessórios, por exemplo) — gere só 1 linha por SISTEMA (rack), ignorando seleções de acessórios que não sejam o conjunto de compressores em si.
- Quantidade_Compressores: conte quantas unidades do mesmo modelo aparecem na seleção total do rack (ex.: "4x 4PES-12Y" = 4).
- HP_Total: SEMPRE "Não informado" — não é informado pelo relatório de nenhum dos dois fabricantes, o usuário preenche manualmente depois no aplicativo.
- Sempre mostre os dados extraídos ao usuário em formato de tabela, para revisão e confirmação, antes de gerar o .xlsx final.
- O arquivo deve ser importado no aplicativo pelo botão de upload na Tela 6 — Rack Paralelo."""

INSTRUCOES_BITZER = f"""Instruções para quem for processar este documento (IA ou humano, sem precisar de contexto prévio):

Objetivo: ler o PDF anexo (relatório de saída do BITZER Software, seleção de compressores semi-herméticos/parafuso em paralelo) e gerar um arquivo .xlsx com esta estrutura (1 aba "Rack"):

Colunas: {COLUNAS}

Onde encontrar cada campo no relatório BITZER Software (seções "Project survey", "Result", "Technical Data", "Selection: Horizontal receivers"):
- Quantidade_Compressores: seção "Project survey" — ex.: "Semi-hermetic Reciprocating Compressors 4x 4PES-12Y" → Quantidade_Compressores = 4.
- Modelo_Compressor: o código do modelo selecionado (ex.: "4PES-12Y").
- Tensao: campo "Power supply" da seção "Selection" (ex.: "380V-3-60Hz").
- Linha_Compressor: o título da família mostrado em "Project survey" (ex.: "Semi-hermetic Reciprocating Compressors").
- Fabricante_Compressor: sempre "Bitzer".
- COP_Compressor: campo "COP/EER" da seção "Result", linha "Total" (ex.: 4,17).
- Capacidade_Compressor_Kcal_H: "Não informado" (o relatório só traz em kW, não por compressor em kcal/h — a conversão fica pro aplicativo/usuário se precisar).
- Carga_Total_Fornecida_Kcal_H: linha "Evaporator capacity" (linha "Total", em kW) da seção "Result" — CONVERTA para kcal/h multiplicando por 860 (ex.: 150,8kW → 129.688 kcal/h).
- Calor_Total_Rejeitado_Kcal_H: linha "Condenser capacity" (linha "Total", em kW) da mesma seção — CONVERTA para kcal/h (×860).
- Potencia_Total_W: linha "Power input" (linha "Total", em kW) — CONVERTA para W (×1000).
- Corrente_Nominal_A: linha "Current (380V)" (ou a tensão configurada), coluna "Total" da seção "Result" — essa é a corrente TOTAL do rack já somada, não confunda com a corrente de 1 compressor.
- Corrente_Maxima_Trabalho_A: campo "Max. operating current" da seção "Technical Data" (⚠ esse é POR COMPRESSOR — some ou deixe unitário mesmo, é isso que o relatório traz).
- Vazao_Massica_Kg_H: linha "Mass flow", coluna "Total" da seção "Result" (kg/h).
- Carga_Oleo_L: campo "Oil charge" da seção "Technical Data" (vem em dm³, equivale a litros) — multiplique pela Quantidade_Compressores se quiser o total do rack, ou deixe unitário e avise o usuário na tabela de revisão.
- Conexao_Descarga: campo "Connection discharge line" da seção "Technical Data" (ex.: "28mm - 1 1/8''").
- Conexao_Succao: campo "Connection suction line" da seção "Technical Data" (ex.: "35mm - 1 3/8''").
- Tanque_Liquido_L: se houver seção "Selection: Horizontal receivers" no relatório, campo "Receiver volume" (ex.: "160,0 dm³" = 160 L). Se não houver essa seção no PDF, "Não informado".
- Carga_Gas_Estimada_Kg: na mesma seção de receiver, campo "Max. refrigerant charge" para o gás do projeto (ex.: linha "R134a" na tabela "Technical Data" do receiver). Se não houver essa seção, "Não informado".

{REGRAS_COMUNS}"""

INSTRUCOES_COPELAND = f"""Instruções para quem for processar este documento (IA ou humano, sem precisar de contexto prévio):

Objetivo: ler o PDF anexo (relatório de saída do Copeland Select / Emerson climate.emerson.com, seleção de compressores) e gerar um arquivo .xlsx com esta estrutura (1 aba "Rack"):

Colunas: {COLUNAS}

Onde encontrar cada campo no relatório Copeland (seções "Project Report", "PERFORMANCE AT SPECIFIED OPERATING POINT"):
- Quantidade_Compressores: procure no relatório quantas unidades do modelo selecionado compõem o rack (ex.: "4 x D3DS-100X" → Quantidade_Compressores = 4). Esse texto pode aparecer perto do campo "Required Capacity", na página de resumo do projeto.
- Modelo_Compressor: o código do modelo (ex.: "D3DS-100X").
- Tensao: o relatório Copeland pode não trazer a tensão explicitamente na página de performance — procure em "Current at 380 V" (aí a tensão é 380V) ou pergunte ao usuário; se não encontrar, "Não informado".
- Linha_Compressor: a família/série do modelo (ex.: prefixo "D3DS" já indica a linha) — se não estiver claro, "Não informado".
- Fabricante_Compressor: sempre "Copeland".
- COP_Compressor: campo "COP" da seção "PERFORMANCE AT SPECIFIED OPERATING POINT" (valor por unidade, ex.: 4.15).
- Capacidade_Compressor_Kcal_H: campo "Cooling Capacity" (kW, por unidade) da mesma seção — CONVERTA para kcal/h (×860).
- Carga_Total_Fornecida_Kcal_H: campo "Refrigeration Capacity" (kW) — geralmente já é o total do rack, não por unidade; confirme pelo valor (se bater com Cooling Capacity × Quantidade_Compressores, é o total). CONVERTA para kcal/h (×860).
- Calor_Total_Rejeitado_Kcal_H: campo "Heating Capacity" (kW, por unidade) — se for por unidade, multiplique pela Quantidade_Compressores antes de converter; CONVERTA o total para kcal/h (×860).
- Potencia_Total_W: campo "Power Input" (kW) da seção de performance — se for total do rack, CONVERTA direto para W (×1000); se for só "Power" por unidade, multiplique pela Quantidade_Compressores antes.
- Corrente_Nominal_A: campo "Current at 380 V" (ou a tensão do projeto) — é por unidade; multiplique pela Quantidade_Compressores para o total do rack, ou deixe unitário e avise na tabela de revisão.
- Corrente_Maxima_Trabalho_A: o relatório Copeland normalmente NÃO traz esse dado — "Não informado".
- Vazao_Massica_Kg_H: campo "Suction Mass Flow" (geralmente em g/s, por unidade) — CONVERTA para kg/h (× 3,6) e multiplique pela Quantidade_Compressores para o total.
- Carga_Oleo_L: o relatório Copeland normalmente NÃO traz esse dado — "Não informado".
- Conexao_Descarga / Conexao_Succao: o relatório Copeland pode trazer numa página de desenho/dimensões do modelo (ex.: "DL 1 1/8 ODS" = conexão de descarga) — extraia se estiver claro, senão "Não informado".
- Tanque_Liquido_L / Carga_Gas_Estimada_Kg: o Copeland Select normalmente não seleciona reservatório de líquido nesse relatório — "Não informado".

{REGRAS_COMUNS}"""


def _titulo(doc, texto, tamanho=16, cor=RGBColor(0x11, 0x18, 0x27)):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(tamanho)
    run.font.color.rgb = cor
    return p


def _gerar(nome_fabricante, instrucoes) -> bytes:
    doc = Document()
    _titulo(doc, f"Documento de Importação — Rack Paralelo ({nome_fabricante})", 18)
    intro = doc.add_paragraph(
        f"Anexe junto com este documento o PDF do relatório de seleção de compressores ({nome_fabricante}) "
        "numa conversa com o Claude, pedindo para gerar a planilha Excel de importação — as instruções "
        "já estão neste documento. Informe também qual Sistema do projeto (Tela 1) esse relatório representa."
    )
    intro.runs[0].font.size = Pt(11)
    doc.add_paragraph()

    _titulo(doc, "Instruções para quem for processar este documento", 13, RGBColor(0xB4, 0x53, 0x09))
    for linha in instrucoes.split("\n"):
        p = doc.add_paragraph(linha or " ")
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def gerar_documento_importacao_rack_bitzer() -> bytes:
    return _gerar("Bitzer", INSTRUCOES_BITZER)


def gerar_documento_importacao_rack_copeland() -> bytes:
    return _gerar("Copeland", INSTRUCOES_COPELAND)
