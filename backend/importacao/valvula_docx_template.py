# -*- coding: utf-8 -*-
"""Gera o "Documento de Importação — Válvulas de Expansão" (.docx). Diferente do fluxo de
Forçador/UC (que colam prints de catálogo), aqui a fonte é o PDF COMPLETO do relatório de
dimensionamento gerado pelo próprio app do fabricante (Danfoss Coolselector2, Full Gauge VEE
Selector ou Carel CPQ) — o usuário anexa esse PDF junto com este documento numa conversa com o
Claude (ou outra IA), que devolve a planilha .xlsx pronta para upload no aplicativo."""
import io
from docx import Document
from docx.shared import Pt, RGBColor

INSTRUCOES_IA = """Instruções para quem for processar este documento (IA ou humano, sem precisar de contexto prévio):

Objetivo: ler o PDF anexo (relatório de dimensionamento de Válvula de Expansão — Danfoss Coolselector2, Full Gauge VEE Selector ou Carel CPQ, qualquer um dos três formatos) e gerar um arquivo .xlsx com esta estrutura (1 aba):

Aba "Valvulas": Identificador, Fabricante, Tipo_Expansao, Modelo_Selecao, Capacidade_Unit_Kcal_h, Orificio, Carga_Abertura_Pct, Conexao_Entrada, Conexao_Saida, Tensao, Tipo_Motor — uma linha por válvula/seleção encontrada no relatório.

O que é cada coluna:
- Identificador: o texto que o usuário digitou no campo "Nome da Aplicação" (Full Gauge), "Nome do Projeto"/nome da seleção (Danfoss) ou nome da câmara/sistema (Carel) dentro do relatório. Formato esperado: código da câmara seguido, sem espaço, de "F" + um número (ex.: "HTA1AabF1", "MTB2CcdF2"). Transcreva EXATAMENTE como está no relatório, sem corrigir nem reformatar.
- Fabricante: "Danfoss", "Fullgauge" ou "Carel", conforme o formato do relatório.
- Tipo_Expansao: "Eletrônica" para válvula eletrônica (EEV/VEE — Full Gauge, Carel), "Termostática" para válvula termostática mecânica (TXV — Danfoss).
- Modelo_Selecao: o MODELO da válvula na nomenclatura do banco de dados do aplicativo — NUNCA o código/part number do fornecedor. Full Gauge: só o modelo, sem o código numérico (ex.: relatório "04336 - SB130T" -> escreva "SB130T"; "04492 - SB520" -> "SB520"). Carel: o modelo (ex.: "E2V 24"). Danfoss: o corpo vem como "TE 2"/"TE 5" etc. — converta pra família do banco conforme o GÁS refrigerante do relatório: R22 -> TEX, R407C -> TEZ, R134a -> TEN, R404A/R507 -> TES; junte família + tamanho SEM espaço (ex.: seleção "TE 5 - 2" com R134a -> Modelo "TEN5"; "T2 - 4" com R134a -> "TEN2"; "TE 5 - 3" com R404A -> "TES5"). NÃO transcreva códigos como 068Z3262/067B3297.
- Capacidade_Unit_Kcal_h: a capacidade nominal/de trabalho da válvula na condição selecionada, em kcal/h (converta se o relatório mostrar em kW ou TR: 1 kW = 860 kcal/h; 1 TR = 3.024 kcal/h — indique a conversão ao usuário na revisão). Se o relatório não trouxer, deixe em branco.
- Orificio: SÓ para válvula termostática (Danfoss) — o NÚMERO do orifício da seleção (o número depois do hífen no tipo selecionado: "TE 5 - 2" -> 2; "T2 - 4" -> 4; "T2 - 0" -> 0). NÃO use o código do fornecedor (068-2007 etc.). Eletrônica não tem orifício: deixe em branco.
- Carga_Abertura_Pct: o percentual de carga/abertura/ponto de trabalho da válvula (campo "Carga [%]" na Danfoss, "Valor da Abertura" na Full Gauge). Só o número, sem o símbolo %. Se o relatório não trouxer esse dado, deixe em branco.
- Conexao_Entrada / Conexao_Saida: diâmetro de conexão de entrada/saída (ex.: 3/8", 1/2"), quando o relatório trouxer essa informação explicitamente (Full Gauge traz "ø Entrada"/"ø Saída"; Danfoss geralmente NÃO traz esse dado nesse relatório — nesse caso deixe em branco, não invente; Carel às vezes traz embutido no texto da descrição do part number, ex.: "16(5/8\")" — extraia se estiver claro, senão deixe em branco).
- Tensao: SÓ para válvula eletrônica — tensão de alimentação do motor da válvula, quando o relatório trouxer (ex.: "12 Vdc <> 10%"). Termostática não tem: deixe em branco.
- Tipo_Motor: SÓ para válvula eletrônica — "Unipolar" ou "Bipolar", quando o relatório trouxer. Termostática não tem: deixe em branco.
- Observação: campos mecânicos que o relatório não trouxer (Conexões/Tensao/Tipo_Motor) serão completados automaticamente pelo banco de dados do aplicativo na importação — por isso não invente: em caso de divergência, o RELATÓRIO tem prioridade, então só preencha o que estiver realmente no relatório.

Regras importantes:
- NUNCA invente dado que não está no relatório. Campo que não aparece fica em branco na planilha — nunca omita a coluna.
- Um mesmo PDF pode conter vários relatórios (uma seleção por câmara/ambiente) — gere uma linha na aba "Valvulas" para CADA seleção encontrada, mesmo que sejam dezenas.
- Quando o relatório mostrar mais de uma OPÇÃO candidata para a mesma seleção (ex.: Danfoss T2-4/T2-5/T2-6), extraia APENAS a opção efetivamente recomendada/destacada (normalmente marcada em verde ou indicada como seleção final) — não gere uma linha por opção descartada.
- Se o Identificador não seguir o padrão esperado (código da câmara + "F" + número) ou não for encontrado no relatório, transcreva o que houver mesmo assim — a conferência de qual câmara/forçador corresponde é feita depois, na prévia do aplicativo.
- Sempre mostre os dados extraídos ao usuário em formato de tabela, para revisão e confirmação, antes de gerar o .xlsx final.
- O arquivo deve ser importado no aplicativo pelo botão de upload na Tela 2 — Câmara Completo (cobre válvulas de câmaras Completo e Simples do mesmo projeto num upload só)."""


def _titulo(doc, texto, tamanho=16, cor=RGBColor(0x11, 0x18, 0x27)):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(tamanho)
    run.font.color.rgb = cor
    return p


def gerar_documento_importacao_valvulas() -> bytes:
    doc = Document()

    _titulo(doc, "Documento de Importação — Válvulas de Expansão", 18)
    intro = doc.add_paragraph(
        "Anexe junto com este documento o PDF do relatório de dimensionamento de válvulas "
        "(Danfoss Coolselector2, Full Gauge VEE Selector ou Carel CPQ) numa conversa com o "
        "Claude, pedindo para gerar a planilha Excel de importação — as instruções já estão "
        "neste documento."
    )
    intro.runs[0].font.size = Pt(11)
    doc.add_paragraph()

    _titulo(doc, "Instruções para quem for processar este documento", 13, RGBColor(0xB4, 0x53, 0x09))
    for linha in INSTRUCOES_IA.split("\n"):
        p = doc.add_paragraph(linha or " ")
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
