# -*- coding: utf-8 -*-
"""Tela D — Cálculo de Consumo Elétrico e Payback, exportação Excel. Uma tabela por sistema
(cenário Projeto/Simples) + bloco de Total do Projeto — mesmo padrão visual das demais
exportações (ver backend/exportacao/_estilo.py e compilacao_export.py)."""
import io
from openpyxl import Workbook
from ._estilo import (celula, barra, cabecalho_tabela, autosize, fmt_br,
                      CENTRO, ESQUERDA, ESQUERDA_LONGO, COR_TITULO, COR_TOTAL, COR_MUTED)

COLUNAS = ["Cenário", "Compressor (kWh/mês)", "Ventiladores (kWh/mês)", "Degelo (kWh/mês)",
           "Portas/Drenos (kWh/mês)", "Iluminação (kWh/mês)", "Total (kWh/mês)", "Custo/mês (R$)"]
_NCOLS = len(COLUNAS)


def _linha_cenario(nome, cenario):
    return [nome, cenario["kwh_compressor"], cenario["kwh_ventiladores"], cenario["kwh_degelo"],
            cenario["kwh_portas_drenos"], cenario["kwh_iluminacao"], cenario["kwh_total_mes"], cenario["custo_mes"]]


def gerar_excel_consumo(projeto, dados: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Consumo Elétrico"
    linha = 1

    celula(ws, linha, 1, f"CONSUMO ELÉTRICO — {projeto.codigo_projeto or ''}", negrito=True,
           tamanho=14, alinhamento=ESQUERDA_LONGO, borda=False)
    linha += 2

    for s in dados["sistemas"]:
        barra(ws, linha, f"SISTEMA {s['sistema_nome']}", _NCOLS)
        linha += 1
        if not s.get("disponivel"):
            celula(ws, linha, 1, s.get("motivo") or "Indisponível.", alinhamento=ESQUERDA_LONGO, borda=False)
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
            linha += 2
            continue
        celula(ws, linha, 1, (
            f"UC considerada: {s.get('modelo_uc') or '—'} ({s.get('hp_uc') or '—'} HP) · "
            f"Carga: {s.get('carga_total_kcal_h') or '—'} kcal/h · "
            f"Capacidade UC: {s.get('capacidade_uc_kcal_h') or '—'} kcal/h · "
            f"Fator de carga: {s.get('fator_carga_pct') or '—'}% · Horas/dia: {s.get('horas_funcionamento') or '—'}"),
            cor_fonte=COR_MUTED, tamanho=9, alinhamento=ESQUERDA_LONGO, borda=False)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
        linha += 1
        cabecalho_tabela(ws, linha, COLUNAS)
        linha += 1
        for nome, cenario in (("Projeto (conforme configurado)", s["projeto"]), ("Simples (sem otimizações)", s["simples"])):
            for c, v in enumerate(_linha_cenario(nome, cenario), start=1):
                celula(ws, linha, c, v, alinhamento=ESQUERDA if c == 1 else CENTRO)
            linha += 1
        payback = s.get("payback_meses")
        celula(ws, linha, 1, (
            f"Economia mensal: R$ {s.get('economia_mes')} · "
            f"Payback: {f'{payback} mes(es)' if payback is not None else '—'}"),
            negrito=True, cor_fundo=COR_TOTAL, alinhamento=ESQUERDA_LONGO)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
        linha += 2

    # Total do Projeto
    barra(ws, linha, "TOTAL DO PROJETO", _NCOLS, cor_fundo=COR_TITULO)
    linha += 1
    cabecalho_tabela(ws, linha, ["Cenário", "Total (kWh/mês)", "Custo/mês (R$)"])
    linha += 1
    t = dados["total"]
    for nome, kwh, custo in (("Projeto", t["kwh_projeto_mes"], t["custo_projeto_mes"]),
                              ("Simples", t["kwh_simples_mes"], t["custo_simples_mes"])):
        celula(ws, linha, 1, nome, alinhamento=ESQUERDA)
        celula(ws, linha, 2, kwh)
        celula(ws, linha, 3, custo)
        linha += 1
    payback_total_txt = f"{t['payback_meses']} mes(es)" if t.get("payback_meses") is not None else "—"
    celula(ws, linha, 1, f"Economia mensal total: R$ {t['economia_mes']} · Payback total: {payback_total_txt}",
           negrito=True, cor_fundo=COR_TOTAL, alinhamento=ESQUERDA_LONGO)
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
    linha += 2

    obs = dados.get("observacao")
    if obs:
        celula(ws, linha, 1, "OBSERVAÇÕES TÉCNICAS", negrito=True, tamanho=12, alinhamento=ESQUERDA_LONGO, borda=False)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
        linha += 1
        for trecho in obs.split("\n"):
            celula(ws, linha, 1, trecho, cor_fonte=COR_MUTED, tamanho=9, alinhamento=ESQUERDA_LONGO, borda=False)
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
            linha += 1

    ws.freeze_panes = "A1"
    autosize(ws, _NCOLS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
