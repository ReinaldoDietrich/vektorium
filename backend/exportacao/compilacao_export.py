"""Exportação da Tela 5 — Compilação Linhas — em Excel, com cabeçalho do projeto (Tela 1)
e a mesma tabela agrupada (Sistema > Linha de Sucção, com subtotais e total geral) já validada
no front-end (frontend/js/tela5.js), incluindo as colunas de Disjuntor por bloco e, em projetos
QD (Quadro Distribuído), a "Alimentações Quadros (QD)"."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# 4º elemento = chave do disjuntor no dict de cada item (igual compilacao.py/tela5.js) — todas
# as 5 cargas têm disjuntor calculado.
ELET_LABELS = [("Ventilação", "vent_w", "vent_a", "disjuntor_ventilacao"),
               ("Res. Portas", "portas_w", "portas_a", "disjuntor_res_porta"),
               ("Res. Drenos", "drenos_w", "drenos_a", "disjuntor_res_dreno"),
               ("Degelo", "degelo_w", "degelo_a", "disjuntor_degelo"),
               ("Ilum. Expositores/ Ilum. Ambiente Câm.", "ilum_w", "ilum_a", "disjuntor_iluminacao")]
# Cor de fundo do cabeçalho de cada bloco — mesma separação visual da tela (bordas), aqui por cor.
BLOCO_CORES = {"Ventilação": "DBEAFE", "Res. Portas": "FDE7CE", "Res. Drenos": "DCFCE7",
               "Degelo": "FEE2E2", "Ilum. Expositores/ Ilum. Ambiente Câm.": "EDE9FE",
               "Alimentações Quadros (QD)": "FEF3C7"}
COLUNAS_BASE = ["Sistema", "Gabinete/Câmara", "Modulação/Área", "Carga Térmica", "Forçadores",
                "Folga Real", "Nº Trocas de Ar", "Válvulas de Expansão"]
TIPO_ORDER = {"expositor": 0, "simplificado": 1, "completo": 2}

# Cópia local de compilacao.py:COLUNAS_RESUMO — só usada como default se o chamador não passar
# colunas_resumo (evita import circular, já que compilacao.py importa as funções deste módulo).
COLUNAS_RESUMO_DEFAULT = [
    ("pot_total_instalada_w", "Pot. Total Instalada (W)", "potencia_total_instalada_w"),
    ("pot_maxima_w", "Pot. Máxima (W)", "potencia_total_w"),
    ("pot_maxima_kva", "Pot. Máxima (kVA)", "potencia_total_kva"),
    ("pot_demandada_w", "Pot. em Operação (W)", "potencia_demandada_w"),
    ("pot_demandada_kva", "Pot. em Operação (kVA)", "potencia_demandada_kva"),
    ("tensao", "Tensão", "tensao"),
    ("cabos", "Cabos", "cabos"),
    ("disjuntor", "Disjuntor Alimentação (Sugerido)", "disjuntor"),
]

# Mesma cópia local, agora pro resumo de Compressão/Condensação (mesmas chaves/ordem do Quadro de
# Linhas, exceto "Pot. Total Instalada" — sem esse conceito aqui). Rótulos levam a sigla da
# corrente-fonte (MCC/RLA), já que aqui os dois valores vêm direto do catálogo.
COLUNAS_RESUMO_COMPRESSAO_DEFAULT = [
    ("pot_maxima_w", "Pot. Máxima - MCC (W)", "potencia_maxima_w"),
    ("pot_maxima_kva", "Pot. Máxima - MCC (kVA)", "potencia_maxima_kva"),
    ("pot_demandada_w", "Pot. em Operação - RLA (W)", "potencia_demandada_w"),
    ("pot_demandada_kva", "Pot. em Operação - RLA (kVA)", "potencia_demandada_kva"),
    ("tensao", "Tensão", "tensao"),
    ("cabos", "Cabos", "cabos"),
    ("disjuntor", "Disjuntor Alimentação (Sugerido)", "disjuntor"),
]

_BORDA = Border(left=Side(style="thin", color="D1D5DB"), right=Side(style="thin", color="D1D5DB"),
                 top=Side(style="thin", color="D1D5DB"), bottom=Side(style="thin", color="D1D5DB"))
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ESQUERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)
# Texto longo (linhas de barra/subtotal/total/observação) — não quebra (a célula é mesclada por
# várias colunas, então tem largura de sobra e não empurra a coluna A pra ficar gigante).
_ESQUERDA_LONGO = Alignment(horizontal="left", vertical="center", wrap_text=False)


def _fmt(n, dec=0):
    if n is None:
        return "—"
    return f"{n:,.{dec}f}".replace(",", "@").replace(".", ",").replace("@", ".")


def _total_resumo(resumo_sistemas, colunas_resumo):
    """Soma por coluna numérica (W/kVA) pra linha 'Alimentação Geral'. Tensão/Cabos/Disjuntor não somam
    (Disjuntor geral agregado vem pronto do backend, não é soma de W/kVA)."""
    return {chave: sum((r.get(campo) or 0) for r in resumo_sistemas)
            for chave, _, campo in colunas_resumo if chave.endswith("_w") or chave.endswith("_kva")}


def _agrupar(itens):
    """Mesma lógica de agrupamento do tela5.js — Sistema > Linha de Sucção, com subtotais."""
    por_sistema = {}
    for c in itens:
        por_sistema.setdefault(c["sistema_nome"], []).append(c)
    grupos = []
    gran_t, gran_e = 0.0, 0.0
    for sistema, lista_sistema in por_sistema.items():
        ref = lista_sistema[0]
        por_ramal = {}
        for c in lista_sistema:
            por_ramal.setdefault(c["linha_succao"], []).append(c)
        ramais = []
        tot_t, tot_e = 0.0, 0.0
        for ramal in sorted(por_ramal):
            lista = sorted(por_ramal[ramal], key=lambda c: TIPO_ORDER.get(c["metodo"], 9))
            sub_t, sub_e = 0.0, 0.0
            for c in lista:
                sub_t += c.get("carga_termica") or 0
                sub_e += sum(c.get(w) or 0 for _, w, _, _ in ELET_LABELS)
            tot_t += sub_t
            tot_e += sub_e
            ramais.append({"ramal": ramal, "itens": lista, "subtotal_carga": sub_t, "subtotal_eletrica": sub_e})
        gran_t += tot_t
        gran_e += tot_e
        grupos.append({"sistema": sistema, "temp_evaporacao": ref.get("temp_evaporacao"),
                        "gas_refrigerante": ref.get("gas_refrigerante"), "ramais": ramais,
                        "total_carga": tot_t, "total_eletrica": tot_e})
    return grupos, gran_t, gran_e


def _cabecalho_linhas(projeto, sistemas):
    linhas = [
        ("Código Projeto", projeto.codigo_projeto or "—"),
        ("Data", projeto.data or "—"),
        ("Cidade de Instalação", projeto.cidade_instalacao or "—"),
        ("Altitude (m)", projeto.altitude_m if projeto.altitude_m is not None else "—"),
        ("Temperatura Ambiente (°C)", projeto.temp_ambiente if projeto.temp_ambiente is not None else "—"),
        ("UR Externa (%)", projeto.ur_externa if projeto.ur_externa is not None else "—"),
        ("Tipo de Comando", projeto.tipo_comando or "—"),
        ("Tensão dos Equipamentos", projeto.tensao_equipamentos or "—"),
        ("Tensão de Comando", projeto.tensao_comando or "—"),
        ("Custo de Energia (R$/kWh)", projeto.custo_energia if projeto.custo_energia is not None else "—"),
        ("Condição do Salão", projeto.condicao_salao or "—"),
    ]
    for s in sistemas:
        detalhe = " · ".join(filter(None, [s.classificacao, s.gas_refrigerante, s.tipo_expansao,
                                            f"{s.temp_evaporacao}°C evap." if s.temp_evaporacao is not None else None,
                                            s.tipo_automacao_linhas, s.automacao_linhas_fabricante]))
        linhas.append((f"Sistema: {s.nome}", detalhe or "—"))
    return linhas


def _celula(ws, linha, coluna, valor, *, negrito=False, cor_fonte=None, cor_fundo=None,
            alinhamento=_CENTRO, tamanho=None, borda=True):
    cel = ws.cell(row=linha, column=coluna, value=valor)
    cel.font = Font(bold=negrito, color=cor_fonte, size=tamanho or 11)
    if cor_fundo:
        cel.fill = PatternFill("solid", fgColor=cor_fundo)
    cel.alignment = alinhamento
    if borda:
        cel.border = _BORDA
    return cel


def _autosize_colunas(ws, ncols, minimo=9, maximo=42):
    # Células mescladas (barras SISTEMA/TOTAL/Subtotal, Total Geral, Observações) ficam de fora do
    # cálculo — senão o texto longo delas, que só vive na coluna A mas se espalha por várias colunas,
    # inflaria a coluna A e desalinharia tudo.
    celulas_mescladas = set()
    for rng in ws.merged_cells.ranges:
        for linha_r in range(rng.min_row, rng.max_row + 1):
            for col_r in range(rng.min_col, rng.max_col + 1):
                celulas_mescladas.add((linha_r, col_r))
    larguras = {}
    for row in ws.iter_rows():
        for cel in row:
            if cel.column > ncols or cel.value is None or (cel.row, cel.column) in celulas_mescladas:
                continue
            tamanho = max(len(t) for t in str(cel.value).split("\n"))
            larguras[cel.column] = max(larguras.get(cel.column, 0), tamanho)
    for col in range(1, ncols + 1):
        letra = ws.cell(row=1, column=col).column_letter
        ws.column_dimensions[letra].width = max(minimo, min(maximo, larguras.get(col, minimo) + 2))


def gerar_excel_compilacao(projeto, sistemas, itens, resumo_sistemas=None, observacao=None, colunas_resumo=None,
                            resumo_compressao=None, colunas_resumo_compressao=None, totais_gerais=None,
                            eh_qd=False, disjuntor_geral_quadro_linhas=None, disjuntor_geral_compressao=None) -> bytes:
    colunas_resumo = colunas_resumo or COLUNAS_RESUMO_DEFAULT
    colunas_resumo_compressao = colunas_resumo_compressao or COLUNAS_RESUMO_COMPRESSAO_DEFAULT
    wb = Workbook()
    ws = wb.active
    ws.title = "Compilação"
    resumo_por_sistema = {r["sistema_nome"]: r for r in (resumo_sistemas or [])}

    ncols_tabela = len(COLUNAS_BASE) + len(ELET_LABELS) * 3 + (2 if eh_qd else 0)
    ncols_total = max(ncols_tabela, len(colunas_resumo) + 1, len(colunas_resumo_compressao) + 1)

    linha = 1
    ws.cell(row=linha, column=1, value="COMPILAÇÃO — CARGA TÉRMICA").font = Font(bold=True, size=14)
    linha += 2
    for label, valor in _cabecalho_linhas(projeto, sistemas):
        ws.cell(row=linha, column=1, value=label).font = Font(bold=True)
        ws.cell(row=linha, column=2, value=valor)
        linha += 1
    linha += 1

    # Cabeçalho em 2 linhas: bloco (mesclado, colorido por carga) + sub-coluna (W/A/Disjuntor) —
    # mesma estrutura visual já usada na tela (frontend/js/tela5.js:t5_buildHeaderRows).
    header_row1 = linha
    header_row2 = linha + 1
    for i, nome_col in enumerate(COLUNAS_BASE, start=1):
        _celula(ws, header_row1, i, nome_col, negrito=True, cor_fonte="FFFFFF", cor_fundo="111827")
        _celula(ws, header_row2, i, None, cor_fundo="111827")
        ws.merge_cells(start_row=header_row1, start_column=i, end_row=header_row2, end_column=i)
    col = len(COLUNAS_BASE) + 1
    for lbl, _, _, _ in ELET_LABELS:
        cor = BLOCO_CORES[lbl]
        _celula(ws, header_row1, col, lbl, negrito=True, cor_fundo=cor)
        ws.merge_cells(start_row=header_row1, start_column=col, end_row=header_row1, end_column=col + 2)
        for j, sub in enumerate(("W", "A", "Disjuntor")):
            _celula(ws, header_row2, col + j, sub, negrito=True, cor_fundo=cor)
        col += 3
    if eh_qd:
        cor = BLOCO_CORES["Alimentações Quadros (QD)"]
        _celula(ws, header_row1, col, "Alimentações Quadros (QD)", negrito=True, cor_fundo=cor)
        ws.merge_cells(start_row=header_row1, start_column=col, end_row=header_row1, end_column=col + 1)
        for j, sub in enumerate(("A", "Disjuntor")):
            _celula(ws, header_row2, col + j, sub, negrito=True, cor_fundo=cor)
        col += 2
    ws.row_dimensions[header_row1].height = 18
    ws.freeze_panes = f"A{header_row2 + 1}"
    linha = header_row2 + 1

    grupos, gran_t, _gran_e = _agrupar(itens)
    gran_total_w = sum(r["potencia_total_w"] for r in resumo_por_sistema.values())
    gran_demandada_w = sum(r["potencia_demandada_w"] for r in resumo_por_sistema.values())
    for g in grupos:
        _celula(ws, linha, 1, f"SISTEMA {g['sistema']} — Evap. {g['temp_evaporacao']}°C — Gás {g['gas_refrigerante']}",
                negrito=True, cor_fonte="FFFFFF", cor_fundo="374151", alinhamento=_ESQUERDA_LONGO)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols_tabela)
        linha += 1
        for r in g["ramais"]:
            _celula(ws, linha, 1, f"Linha de Sucção {r['ramal']}", negrito=True, cor_fonte="6B7280", alinhamento=_ESQUERDA, borda=False)
            linha += 1
            for c in r["itens"]:
                nome = c["nome"] if c["metodo"] != "expositor" else c["nome"] + (f" - {c['setor']}" if c.get("setor") else "")
                _celula(ws, linha, 1, c["codigo"], negrito=True, alinhamento=_ESQUERDA)
                _celula(ws, linha, 2, nome, alinhamento=_ESQUERDA)
                _celula(ws, linha, 3, c["modulacao_area"])
                _celula(ws, linha, 4, c.get("carga_termica"))
                _celula(ws, linha, 5, c["forcadores"], alinhamento=_ESQUERDA)
                _celula(ws, linha, 6, c["folga_real"])
                _celula(ws, linha, 7, c["trocas_ar"])
                _celula(ws, linha, 8, c["valvulas"], alinhamento=_ESQUERDA)
                col = len(COLUNAS_BASE) + 1
                for _, w, a, d in ELET_LABELS:
                    _celula(ws, linha, col, c.get(w))
                    _celula(ws, linha, col + 1, c.get(a))
                    _celula(ws, linha, col + 2, c.get(d) or "—", alinhamento=_ESQUERDA)
                    col += 3
                if eh_qd:
                    _celula(ws, linha, col, c.get("alimentacao_quadro_a"))
                    _celula(ws, linha, col + 1, c.get("disjuntor_alimentacao_quadro") or "—", alinhamento=_ESQUERDA)
                    col += 2
                linha += 1
            _celula(ws, linha, 1, f"Subtotal Linha {r['ramal']}: {_fmt(r['subtotal_carga'])} kcal/h · {_fmt(r['subtotal_eletrica'])} W",
                    negrito=True, cor_fundo="F3F4F6", alinhamento=_ESQUERDA_LONGO, borda=False)
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols_tabela)
            linha += 1
        rs = resumo_por_sistema.get(g["sistema"], {})
        _celula(ws, linha, 1, f"TOTAL SISTEMA {g['sistema']}: {_fmt(g['total_carga'])} kcal/h · "
                f"Potência Máxima Instalada {_fmt(rs.get('potencia_total_w'))} W · "
                f"Potência em Operação {_fmt(rs.get('potencia_demandada_w'))} W",
                negrito=True, cor_fonte="FFFFFF", cor_fundo="111827", alinhamento=_ESQUERDA_LONGO, borda=False)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols_tabela)
        linha += 1
    _celula(ws, linha, 1, f"TOTAL GERAL DO PROJETO: {_fmt(gran_t)} kcal/h · "
            f"Potência Máxima Instalada {_fmt(gran_total_w)} W · Potência em Operação {_fmt(gran_demandada_w)} W",
            negrito=True, cor_fonte="FFFFFF", cor_fundo="000000", alinhamento=_ESQUERDA_LONGO, borda=False, tamanho=12)
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols_tabela)
    linha += 2

    def _tabela_resumo(titulo, linhas_dados, rotulo_col1, colunas_def, totais, rotulo_total):
        nonlocal linha
        _celula(ws, linha, 1, titulo, negrito=True, tamanho=12, alinhamento=_ESQUERDA_LONGO, borda=False)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=len(colunas_def) + 1)
        linha += 1
        _celula(ws, linha, 1, rotulo_col1, negrito=True, cor_fonte="FFFFFF", cor_fundo="374151")
        for c, (_, rotulo, _) in enumerate(colunas_def, start=2):
            _celula(ws, linha, c, rotulo, negrito=True, cor_fonte="FFFFFF", cor_fundo="374151")
        linha += 1
        for rotulo_linha, dados in linhas_dados:
            _celula(ws, linha, 1, rotulo_linha, alinhamento=_ESQUERDA)
            for c, (chave, _, campo) in enumerate(colunas_def, start=2):
                valor = dados.get(campo)
                if chave.endswith("_kva") and valor is not None:
                    valor = round(valor, 2)
                _celula(ws, linha, c, valor)
            linha += 1
        _celula(ws, linha, 1, rotulo_total, negrito=True, cor_fundo="F3F4F6", alinhamento=_ESQUERDA)
        for c, (chave, _, campo) in enumerate(colunas_def, start=2):
            valor = totais.get(chave)
            if chave in totais and chave.endswith("_kva") and valor is not None:
                valor = round(valor, 2)
            _celula(ws, linha, c, valor, negrito=True, cor_fundo="F3F4F6")
        linha += 2

    if resumo_sistemas:
        totais = _total_resumo(resumo_sistemas, colunas_resumo)
        totais["disjuntor"] = disjuntor_geral_quadro_linhas
        _tabela_resumo("RESUMO DE POTÊNCIA POR SISTEMA — QUADROS DE LINHAS",
                       [(f"Quadro de Linhas - {r['sistema_nome']}", r) for r in resumo_sistemas],
                       "Sistema", colunas_resumo, totais, "Alimentação Geral - Quadro de Linhas")

    if resumo_compressao:
        totais_comp = _total_resumo(resumo_compressao, colunas_resumo_compressao)
        totais_comp["disjuntor"] = disjuntor_geral_compressao
        _tabela_resumo("RESUMO DE POTÊNCIA POR SISTEMA — COMPRESSÃO E CONDENSAÇÃO",
                       [(f"{r['sistema_nome']} — {r['equipamento']}", r) for r in resumo_compressao],
                       "Sistema / Equipamento", colunas_resumo_compressao, totais_comp, "Alimentação Geral - Compressão e Condensação")

    if totais_gerais:
        # Barra horizontal única (mesmo padrão da tela — frontend/js/tela5.js:t5_tabelaTotalGeral),
        # não linhas verticais. Respeita o filtro de colunas: cada métrica (W/kVA de Máxima e
        # Operação) só entra se a coluna correspondente estiver selecionada em algum dos 2 resumos.
        chaves_sel = {c[0] for c in colunas_resumo} | {c[0] for c in colunas_resumo_compressao}
        maxima = []
        if "pot_maxima_w" in chaves_sel:
            maxima.append(f"{_fmt(totais_gerais.get('potencia_geral_maxima_w'))} W")
        if "pot_maxima_kva" in chaves_sel:
            maxima.append(f"({_fmt(totais_gerais.get('potencia_geral_maxima_kva'), 2)} kVA)")
        operacao = []
        if "pot_demandada_w" in chaves_sel:
            operacao.append(f"{_fmt(totais_gerais.get('potencia_geral_demandada_w'))} W")
        if "pot_demandada_kva" in chaves_sel:
            operacao.append(f"({_fmt(totais_gerais.get('potencia_geral_demandada_kva'), 2)} kVA)")
        partes = []
        if maxima:
            partes.append("Pot. Máxima: " + " ".join(maxima))
        if operacao:
            partes.append("Pot. em Operação: " + " ".join(operacao))
        texto = "TOTAL GERAL — QUADRO DE LINHAS + COMPRESSÃO E CONDENSAÇÃO"
        if partes:
            texto += ": " + " · ".join(partes)
        _celula(ws, linha, 1, texto, negrito=True, cor_fonte="FFFFFF", cor_fundo="111827",
                alinhamento=_ESQUERDA_LONGO, borda=False)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols_total)
        linha += 2

    if observacao:
        _celula(ws, linha, 1, "OBSERVAÇÕES TÉCNICAS", negrito=True, tamanho=12, alinhamento=_ESQUERDA_LONGO, borda=False)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols_total)
        linha += 1
        for trecho in observacao.split("\n"):
            _celula(ws, linha, 1, trecho, cor_fonte="6B7280", tamanho=9, alinhamento=_ESQUERDA_LONGO, borda=False)
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols_total)
            linha += 1

    _autosize_colunas(ws, ncols_total)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
