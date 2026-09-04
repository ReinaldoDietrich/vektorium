# -*- coding: utf-8 -*-
"""Estilo compartilhado das exportações Excel do app — padrão único adotado a partir da Tela 5
(Compilação): bordas em todas as células, cabeçalho colorido/centralizado, texto centralizado,
largura de coluna ajustada ao conteúdo (não fixa/chutada), congelamento do cabeçalho e linhas de
texto longo (barras/observações) mescladas pela largura da tabela. Usado por todas as exportações
de relatório de projeto (Telas 7/8/9) e catálogo (Telas A/B/C) — ver compilacao_export.py para o
mesmo padrão embutido (mantido lá por já estar testado)."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Paleta — mesma da tela/Tela 5.
COR_TITULO = "111827"       # barra preta-azulada (títulos de bloco/sistema)
COR_CABECALHO = "374151"    # cinza-escuro (cabeçalho de tabela)
COR_TOTAL = "F3F4F6"        # cinza-claro (linhas de total/alimentação geral)
COR_PRETO = "000000"        # total geral do projeto
COR_MUTED = "6B7280"        # texto secundário (observações)

_THIN = Side(style="thin", color="D1D5DB")
BORDA = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ESQUERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)
# Texto longo (barras/subtotais/observações) — não quebra: a célula é mesclada por várias colunas,
# então tem largura de sobra e não empurra a coluna 1 pra ficar gigante.
ESQUERDA_LONGO = Alignment(horizontal="left", vertical="center", wrap_text=False)


def fmt_br(n, dec=0):
    """Número no formato brasileiro (1.234,5). None -> '—'."""
    if n is None:
        return "—"
    return f"{n:,.{dec}f}".replace(",", "@").replace(".", ",").replace("@", ".")


def celula(ws, linha, coluna, valor, *, negrito=False, cor_fonte=None, cor_fundo=None,
           alinhamento=None, tamanho=None, borda=True):
    """Escreve uma célula já com o estilo padrão. `alinhamento` default = CENTRO."""
    cel = ws.cell(row=linha, column=coluna, value=valor)
    cel.font = Font(bold=negrito, color=cor_fonte, size=tamanho or 11)
    if cor_fundo:
        cel.fill = PatternFill("solid", fgColor=cor_fundo)
    cel.alignment = alinhamento or CENTRO
    if borda:
        cel.border = BORDA
    return cel


def barra(ws, linha, texto, ncols, *, cor_fundo=COR_TITULO, cor_fonte="FFFFFF", tamanho=11):
    """Linha-título de largura total (bloco/sistema/total), mesclada por `ncols` colunas."""
    celula(ws, linha, 1, texto, negrito=True, cor_fonte=cor_fonte, cor_fundo=cor_fundo,
           alinhamento=ESQUERDA_LONGO, borda=False, tamanho=tamanho)
    if ncols > 1:
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols)


def cabecalho_tabela(ws, linha, titulos, *, cor_fundo=COR_CABECALHO):
    """Escreve a linha de cabeçalho de uma tabela (colorida, centralizada, com borda)."""
    for c, titulo in enumerate(titulos, start=1):
        celula(ws, linha, c, titulo, negrito=True, cor_fonte="FFFFFF", cor_fundo=cor_fundo)


def autosize(ws, ncols, *, minimo=9, maximo=42):
    """Ajusta a largura de cada coluna ao maior conteúdo dela — ignorando células mescladas (o
    texto longo das barras/observações vive só na coluna 1 mas se espalha por várias colunas, e
    inflaria a coluna 1 se contado)."""
    mescladas = set()
    for rng in ws.merged_cells.ranges:
        for lr in range(rng.min_row, rng.max_row + 1):
            for cr in range(rng.min_col, rng.max_col + 1):
                mescladas.add((lr, cr))
    larguras = {}
    for row in ws.iter_rows():
        for cel in row:
            if cel.column > ncols or cel.value is None or (cel.row, cel.column) in mescladas:
                continue
            tam = max(len(t) for t in str(cel.value).split("\n"))
            larguras[cel.column] = max(larguras.get(cel.column, 0), tam)
    for col in range(1, ncols + 1):
        letra = get_column_letter(col)
        ws.column_dimensions[letra].width = max(minimo, min(maximo, larguras.get(col, minimo) + 2))
