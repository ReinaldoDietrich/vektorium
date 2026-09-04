# -*- coding: utf-8 -*-
"""Composição de Preço (Tela 10 fundida) — exportação Excel. Dados do Projeto + um bloco por
seção (Equipamentos/Materiais/Painéis/Mão de Obra/Serviços/Fretes), mesmo padrão visual das
demais exportações (ver backend/exportacao/_estilo.py) — sucede tela10_export.py."""
import io
from openpyxl import Workbook
from ._estilo import (celula, barra, cabecalho_tabela, autosize, fmt_br,
                      CENTRO, ESQUERDA, ESQUERDA_LONGO, COR_MUTED, COR_TOTAL, COR_PRETO)

COLUNAS = ["Item", "Descrição", "Fabricante", "Centro de Custo", "Fator", "Qtd.",
           "Custo Unit.", "Custo Total", "Margem", "Impostos", "Comissão",
           "Vlr. Venda", "C/ Negociação", "Observação"]
_NCOLS = len(COLUNAS)


def _linha_item(item, i):
    return [i, item.get("descricao") or "—", item.get("fabricante") or "—",
            item.get("centro_custo_codigo") or "—", item.get("fator_codigo") or "—",
            item.get("quantidade") or 0,
            fmt_br(item.get("custo_unitario"), 2), fmt_br(item.get("custo_total"), 2),
            fmt_br(item.get("margem_contribuicao"), 2), fmt_br(item.get("impostos"), 2),
            fmt_br(item.get("comissoes"), 2), fmt_br(item.get("valor_total_venda"), 2),
            fmt_br(item.get("valor_venda_negociacao"), 2), item.get("observacao") or ""]


def _bloco_lista(ws, linha, titulo, itens):
    total = sum((it.get("valor_venda_negociacao") or 0) for it in itens)
    barra(ws, linha, f"{titulo}" + (f" — C/ Negociação: R$ {fmt_br(total, 2)}" if total else ""), _NCOLS)
    linha += 1
    cabecalho_tabela(ws, linha, COLUNAS)
    linha += 1
    if not itens:
        celula(ws, linha, 1, "Nenhum item.", alinhamento=ESQUERDA_LONGO, borda=False, cor_fonte=COR_MUTED)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
        linha += 1
    for i, item in enumerate(itens, start=1):
        for c, v in enumerate(_linha_item(item, i), start=1):
            celula(ws, linha, c, v, alinhamento=ESQUERDA if c in (2, 14) else CENTRO)
        linha += 1
    return linha + 1


def gerar_excel_composicao_preco(projeto, composicao: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Composição de Preço"
    linha = 1

    celula(ws, linha, 1, f"RESUMO DE EQUIPAMENTOS, MATERIAIS E COMPOSIÇÃO DE PREÇO — {projeto.codigo_projeto or ''}",
           negrito=True, tamanho=14, alinhamento=ESQUERDA_LONGO, borda=False)
    linha += 2

    barra(ws, linha, "DADOS DO PROJETO", _NCOLS)
    linha += 1
    cidade = getattr(projeto, 'cidade_instalacao', '') or ''
    estado = getattr(projeto, 'estado_uf', '') or ''
    campos = [("Código Projeto", projeto.codigo_projeto),
              ("Cidade/Estado da Instalação", f"{cidade}/{estado}" if (cidade or estado) else "—"),
              ("Tensão Elétrica de Comando", projeto.tensao_comando),
              ("Tensão Elétrica da Instalação", projeto.tensao_equipamentos)]
    for rotulo, valor in campos:
        celula(ws, linha, 1, rotulo, negrito=True, alinhamento=ESQUERDA)
        celula(ws, linha, 2, valor or "—", alinhamento=ESQUERDA)
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=_NCOLS)
        linha += 1
    linha += 1

    for bloco, itens in composicao["blocos"].items():
        linha = _bloco_lista(ws, linha, bloco.upper(), itens)

    ws.freeze_panes = "A1"
    autosize(ws, _NCOLS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_COLUNAS_LISTA_MATERIAIS = ["Nº/Item", "Descrição", "Fabricante", "Centro de Custo", "Unid.", "Quantidade"]


def gerar_excel_lista_materiais(projeto, linhas: list[dict]) -> bytes:
    """"Lista de Materiais e Equipamentos" — versão simples (sem valores financeiros), pra pedido
    de orçamento a fornecedores."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lista de Materiais"
    linha_n = 1

    celula(ws, linha_n, 1, f"LISTA DE MATERIAIS E EQUIPAMENTOS — {projeto.codigo_projeto or ''}",
           negrito=True, tamanho=14, alinhamento=ESQUERDA_LONGO, borda=False)
    linha_n += 1
    cidade = getattr(projeto, 'cidade_instalacao', '') or ''
    estado = getattr(projeto, 'estado_uf', '') or ''
    if cidade or estado:
        celula(ws, linha_n, 1, f"Cidade/Estado: {cidade}/{estado}",
               alinhamento=ESQUERDA_LONGO, borda=False)
        ws.merge_cells(start_row=linha_n, start_column=1, end_row=linha_n, end_column=len(_COLUNAS_LISTA_MATERIAIS))
        linha_n += 1
    linha_n += 1

    cabecalho_tabela(ws, linha_n, _COLUNAS_LISTA_MATERIAIS)
    linha_n += 1
    if not linhas:
        celula(ws, linha_n, 1, "Nenhum item.", alinhamento=ESQUERDA_LONGO, borda=False, cor_fonte=COR_MUTED)
        ws.merge_cells(start_row=linha_n, start_column=1, end_row=linha_n, end_column=len(_COLUNAS_LISTA_MATERIAIS))
        linha_n += 1
    for it in linhas:
        valores = [it["numero"], it.get("descricao") or "—", it.get("fabricante") or "—",
                   it.get("centro_custo") or "—", it.get("unidade") or "—", it.get("quantidade") or 0]
        for c, v in enumerate(valores, start=1):
            celula(ws, linha_n, c, v, alinhamento=ESQUERDA if c == 2 else CENTRO)
        linha_n += 1

    ws.freeze_panes = "A1"
    autosize(ws, len(_COLUNAS_LISTA_MATERIAIS))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_COLUNAS_RESUMO_BLOCO = ["Centro de Custo", "Custo Total", "Impostos", "Comissões",
                          "Margem Contribuição", "Vlr. Venda", "C/ Negociação"]
_CAMPOS_RESUMO_BLOCO = ("custo_total", "impostos", "comissoes", "margem_contribuicao",
                         "valor_total_venda", "valor_venda_negociacao")


def _linha_resumo_bloco(rotulo, valores, *, negrito=False, cor_fundo=None):
    return ([rotulo] + [fmt_br(valores.get(c), 2) for c in _CAMPOS_RESUMO_BLOCO], negrito, cor_fundo)


def gerar_excel_resumo_orcamento(projeto, resumo: dict, orcamento: dict, dre: dict) -> bytes:
    """Resumo por Bloco + Tabela de Orçamento + DRE (Tela 10) — 3 abas, usadas pelo "Exportar
    Tudo" (Tela 1, aprovado 2026-08-12). `resumo`/`orcamento`/`dre` = mesmos dicts retornados por
    resumo_por_bloco/tabela_orcamento/dre_projeto (composicao_preco.py) — nunca recalcula aqui,
    só formata o que a Tela 10 já mostra."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumo por Bloco"
    ncols = len(_COLUNAS_RESUMO_BLOCO)
    linha = 1
    celula(ws, linha, 1, f"RESUMO POR BLOCO — {projeto.codigo_projeto or ''}", negrito=True,
           tamanho=14, alinhamento=ESQUERDA_LONGO, borda=False)
    linha += 2
    for bloco in resumo["blocos"]:
        barra(ws, linha, bloco["bloco"].upper(), ncols)
        linha += 1
        cabecalho_tabela(ws, linha, _COLUNAS_RESUMO_BLOCO)
        linha += 1
        if not bloco["centros_custo"]:
            celula(ws, linha, 1, "Nenhum lançamento.", alinhamento=ESQUERDA_LONGO, borda=False, cor_fonte=COR_MUTED)
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols)
            linha += 1
        for cc in bloco["centros_custo"]:
            rotulo = cc["descricao"] or cc.get("centro_custo") or "Sem Centro de Custo"
            valores, _, _ = _linha_resumo_bloco(rotulo, cc)
            for c, v in enumerate(valores, start=1):
                celula(ws, linha, c, v, alinhamento=ESQUERDA if c == 1 else CENTRO)
            linha += 1
        valores, _, _ = _linha_resumo_bloco(f"Subtotal {bloco['bloco']}", bloco["subtotal"])
        for c, v in enumerate(valores, start=1):
            celula(ws, linha, c, v, negrito=True, cor_fundo=COR_TOTAL, alinhamento=ESQUERDA if c == 1 else CENTRO)
        linha += 2
    valores, _, _ = _linha_resumo_bloco("TOTAL GERAL", resumo["total_geral"])
    for c, v in enumerate(valores, start=1):
        celula(ws, linha, c, v, negrito=True, cor_fonte="FFFFFF", cor_fundo=COR_PRETO,
               alinhamento=ESQUERDA if c == 1 else CENTRO)
    ws.freeze_panes = "A1"
    autosize(ws, ncols)

    ws2 = wb.create_sheet("Tabela de Orçamento")
    colunas_orc = ["Descrição", "Centro de Custo", "Unid.", "Quantidade", "Valor (R$)"]
    ncols2 = len(colunas_orc)
    linha = 1
    celula(ws2, linha, 1, f"TABELA DE ORÇAMENTO — {projeto.codigo_projeto or ''}", negrito=True,
           tamanho=14, alinhamento=ESQUERDA_LONGO, borda=False)
    linha += 2
    for bloco in orcamento["ordem"]:
        dados_bloco = orcamento["blocos"][bloco]
        barra(ws2, linha, f"{bloco.upper()} — R$ {fmt_br(dados_bloco['total'], 2)}", ncols2)
        linha += 1
        cabecalho_tabela(ws2, linha, colunas_orc)
        linha += 1
        if not dados_bloco["itens"] and not dados_bloco["extra_por_cc"]:
            celula(ws2, linha, 1, "Nenhum item.", alinhamento=ESQUERDA_LONGO, borda=False, cor_fonte=COR_MUTED)
            ws2.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols2)
            linha += 1
        for it in dados_bloco["itens"]:
            valores = [it["descricao"], it.get("cc") or "—", it["unidade"], it["quantidade"], fmt_br(it["valor"], 2)]
            for c, v in enumerate(valores, start=1):
                celula(ws2, linha, c, v, alinhamento=ESQUERDA if c == 1 else CENTRO)
            linha += 1
        for chave_cc, valor in dados_bloco["extra_por_cc"].items():
            valores = [f"(Comissão agrupada) {chave_cc}", chave_cc, "-", "-", fmt_br(valor, 2)]
            for c, v in enumerate(valores, start=1):
                celula(ws2, linha, c, v, cor_fonte=COR_MUTED, alinhamento=ESQUERDA if c == 1 else CENTRO)
            linha += 1
        linha += 1
    celula(ws2, linha, 1, "VALOR TOTAL DA PROPOSTA", negrito=True, cor_fonte="FFFFFF", cor_fundo=COR_PRETO,
           alinhamento=ESQUERDA_LONGO, borda=False)
    ws2.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncols2 - 1)
    celula(ws2, linha, ncols2, fmt_br(orcamento["total_geral"], 2), negrito=True, cor_fonte="FFFFFF", cor_fundo=COR_PRETO)
    ws2.freeze_panes = "A1"
    autosize(ws2, ncols2)

    ws3 = wb.create_sheet("DRE")
    linha = 1
    celula(ws3, linha, 1, f"DRE DO PROJETO — {projeto.codigo_projeto or ''}", negrito=True,
           tamanho=14, alinhamento=ESQUERDA_LONGO, borda=False)
    linha += 2
    linhas_dre = [
        ("Receita de Vendas", dre["receita_vendas"], False),
        ("Receita de Comissões de Indicação de Negócio", dre["receita_comissoes"], False),
        ("Receita Bruta Total", dre["receita_bruta"], True),
        ("(-) Impostos sobre Vendas", -dre["impostos_vendas"], False),
        ("(-) Impostos sobre Comissões", -dre["impostos_comissoes"], False),
        ("(-) Impostos", -dre["impostos"], True),
        ("(=) Receita Líquida", dre["receita_liquida"], True),
        ("(-) Custo dos Materiais/Serviços", -dre["custo"], False),
        ("(=) Lucro Bruto", dre["lucro_bruto"], True),
        ("(-) Comissões (reais, por vendedor)", -dre["comissoes"], False),
        ("(=) Resultado do Projeto", dre["resultado"], True),
    ]
    for rotulo, valor, negrito in linhas_dre:
        celula(ws3, linha, 1, rotulo, negrito=negrito, alinhamento=ESQUERDA,
               cor_fundo=COR_TOTAL if negrito else None)
        celula(ws3, linha, 2, f"R$ {fmt_br(valor, 2)}", negrito=negrito, cor_fundo=COR_TOTAL if negrito else None)
        linha += 1
    if dre.get("margem_liquida_pct") is not None:
        celula(ws3, linha, 1, "Margem líquida", alinhamento=ESQUERDA, cor_fonte=COR_MUTED, borda=False)
        celula(ws3, linha, 2, f"{dre['margem_liquida_pct'] * 100:.1f}%".replace(".", ","), cor_fonte=COR_MUTED, borda=False)
    ws3.freeze_panes = "A1"
    autosize(ws3, 2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
