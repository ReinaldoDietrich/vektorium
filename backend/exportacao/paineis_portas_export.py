# -*- coding: utf-8 -*-
"""Tela 7 — Painéis Térmicos e Portas, exportação Excel do Resumo (4 modos: total/painel_portas/
painel_portas_geral/camara) — mesmo padrão visual das demais exportações (ver
backend/exportacao/_estilo.py). PDF fica a cargo da impressão via navegador (frontend)."""
import io
from openpyxl import Workbook
from ._estilo import celula, barra, cabecalho_tabela, autosize, CENTRO, ESQUERDA, ESQUERDA_LONGO, COR_CABECALHO

MODO_TITULO = {
    "total": "Total",
    "painel_portas": "Por Id. Painel + Portas por Câmara",
    "painel_portas_geral": "Por Id. Painel + Portas Lista Geral",
    "camara": "Por Câmara",
}
_NCOLS = 6  # maior tabela = Portas (Descrição, Id., Unid., Quantidade, Tensão, Observação)


def gerar_excel_paineis_portas(projeto, resumo: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Painéis e Portas"
    linha = 1

    def titulo_bloco(texto):
        nonlocal linha
        barra(ws, linha, texto, _NCOLS)
        linha += 1

    def tabela_itens(linhas, titulo, com_area=False):
        nonlocal linha
        if not linhas:
            return
        titulo_bloco(titulo)
        cabecalho = ["Item", "Unid.", "Quantidade"] + (["Área (m²)"] if com_area else [])
        cabecalho_tabela(ws, linha, cabecalho)
        linha += 1
        for it in linhas:
            celula(ws, linha, 1, it["item"], alinhamento=ESQUERDA)
            celula(ws, linha, 2, it["unidade"])
            celula(ws, linha, 3, it["quantidade"])
            if com_area:
                celula(ws, linha, 4, it.get("area_m2"))
            linha += 1
        linha += 1

    def tabela_portas(portas, titulo=None):
        nonlocal linha
        if titulo:
            titulo_bloco(titulo)
        if not portas:
            celula(ws, linha, 1, "Nenhuma porta lançada.", alinhamento=ESQUERDA_LONGO, borda=False)
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
            linha += 2
            return
        cabecalho_tabela(ws, linha, ["Descrição", "Id.", "Unid.", "Quantidade", "Tensão", "Observação"])
        linha += 1
        for p in portas:
            celula(ws, linha, 1, p["descricao"], alinhamento=ESQUERDA)
            celula(ws, linha, 2, p["id_porta"])
            celula(ws, linha, 3, p["unidade"])
            celula(ws, linha, 4, p["quantidade"])
            celula(ws, linha, 5, p.get("tensao") or "")
            celula(ws, linha, 6, p.get("observacoes") or "", alinhamento=ESQUERDA)
            linha += 1
        linha += 1

    celula(ws, linha, 1, f"PAINÉIS E PORTAS — {MODO_TITULO.get(resumo['modo'], resumo['modo'])}",
           negrito=True, tamanho=14, alinhamento=ESQUERDA_LONGO, borda=False)
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
    linha += 1
    celula(ws, linha, 1, f"Projeto: {projeto.codigo_projeto or ''} — {projeto.cliente or ''}",
           alinhamento=ESQUERDA_LONGO, borda=False)
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
    linha += 1
    cidade = getattr(projeto, 'cidade_instalacao', '') or ''
    estado = getattr(projeto, 'estado_uf', '') or ''
    if cidade or estado:
        celula(ws, linha, 1, f"Cidade/Estado: {cidade}/{estado}",
               alinhamento=ESQUERDA_LONGO, borda=False)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
        linha += 1
    linha += 1

    modo = resumo["modo"]
    if modo == "total":
        tabela_itens(resumo["paineis_parede_teto"], "Painéis — Parede e Teto")
        tabela_itens(resumo["isolamento_piso"], "Isolamento de Piso")
        tabela_portas(resumo["portas"], "Portas")

    elif modo in ("painel_portas", "painel_portas_geral"):
        tabela_itens(resumo["paineis_parede_teto"], "Painéis — por Id.", com_area=True)
        tabela_itens(resumo["isolamento_piso"], "Isolamento de Piso")
        if modo == "painel_portas":
            titulo_bloco("Portas por Câmara")
            for c in resumo["portas_por_camara"]:
                titulo_bloco(f"{c['id_planta'] or '—'} — {c['camara_nome']}")
                tabela_portas(c["portas"])
        else:
            tabela_portas(resumo["portas"], "Portas — Lista Geral")

    elif modo == "camara":
        for c in resumo["camaras"]:
            titulo_bloco(f"{c['id_planta'] or '—'} — {c['camara_nome']}")
            tabela_itens(c["paineis_parede_teto"], "Painéis — Parede e Teto")
            tabela_itens(c["isolamento_piso"], "Isolamento de Piso")
            tabela_portas(c["portas"], "Portas")
            linha += 1

    ws.freeze_panes = "A1"
    autosize(ws, _NCOLS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
