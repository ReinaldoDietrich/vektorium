# -*- coding: utf-8 -*-
"""Helper pra combinar múltiplas planilhas Excel (cada uma já gerada por sua função própria de
exportação, SEM modificar nenhuma delas) numa única pasta de trabalho com uma aba por origem —
usado pelo botão "Exportar Tudo" (Tela 1, aprovado 2026-08-12). Reabre cada .xlsx já gerado
(os mesmos bytes que o download individual de cada tela já produz hoje) e copia célula a célula
(valor, formatação, largura de coluna, altura de linha, células mescladas, congelamento de
painel) pra dentro do workbook combinado — nenhuma das 6 funções de exportação individuais
precisa saber lidar com um Workbook compartilhado nem é alterada."""
import io
from copy import copy
import openpyxl


def _copiar_aba(ws_origem, wb_destino, titulo):
    ws = wb_destino.create_sheet(title=titulo[:31])  # limite de 31 caracteres do Excel
    for linha in ws_origem.iter_rows():
        for cel in linha:
            nova = ws.cell(row=cel.row, column=cel.column, value=cel.value)
            if cel.has_style:
                nova.font = copy(cel.font)
                nova.fill = copy(cel.fill)
                nova.border = copy(cel.border)
                nova.alignment = copy(cel.alignment)
                nova.number_format = cel.number_format
    for faixa in ws_origem.merged_cells.ranges:
        ws.merge_cells(str(faixa))
    for letra, dim in ws_origem.column_dimensions.items():
        if dim.width:
            ws.column_dimensions[letra].width = dim.width
    for idx, dim in ws_origem.row_dimensions.items():
        if dim.height:
            ws.row_dimensions[idx].height = dim.height
    if ws_origem.freeze_panes:
        ws.freeze_panes = ws_origem.freeze_panes
    return ws


def adicionar_workbook(wb_destino, conteudo_bytes: bytes, titulo: str):
    """`conteudo_bytes` = arquivo .xlsx já pronto (o mesmo que o download individual da tela
    gera); copia a(s) aba(s) dele pra dentro de wb_destino. Workbook de 1 aba só vira 1 aba
    nomeada `titulo`; de várias abas, cada uma vira "`titulo` - `nome original`"."""
    wb_origem = openpyxl.load_workbook(io.BytesIO(conteudo_bytes))
    if len(wb_origem.sheetnames) == 1:
        _copiar_aba(wb_origem.active, wb_destino, titulo)
    else:
        for nome in wb_origem.sheetnames:
            _copiar_aba(wb_origem[nome], wb_destino, f"{titulo} - {nome}")
