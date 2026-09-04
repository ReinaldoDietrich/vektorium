# -*- coding: utf-8 -*-
"""Tela 11 — Estudo Luminotécnico, exportação Excel. Mesma tabela da tela (linhas por câmara +
resumo somatório por Modelo + notas técnicas), mesmo padrão visual das demais exportações (ver
backend/exportacao/_estilo.py e consumo_export.py) -- aprovado 2026-08-10."""
import io
from openpyxl import Workbook
from ._estilo import (celula, barra, cabecalho_tabela, autosize,
                      CENTRO, ESQUERDA, ESQUERDA_LONGO, COR_TITULO, COR_TOTAL, COR_MUTED)

# Mesmas colunas/rótulos de T18_COLUNAS (frontend/js/telaLuminotecnico.js) -- "modelo_luminaria"
# saiu da tabela principal e virou chave do resumo (aprovado 2026-08-10).
COLUNAS = [
    ("seq", "Seq."), ("linha_succao", "Linha de Sucção"), ("ambiente", "Ambiente"),
    ("comprimento", "Comprimento (m)"), ("largura", "Largura (m)"), ("area", "Área (m²)"),
    ("altura", "Altura (m)"), ("plano_calculo", "Plano de Cálculo (m)"), ("altura_util", "Altura Útil (m)"),
    ("indice_ambiente_k", "Índice do Ambiente (K)"), ("refletancia_teto", "Refletância Teto"),
    ("refletancia_parede", "Refletância Parede"), ("refletancia_piso", "Refletância Piso"),
    ("fator_manutencao", "Fator de Manutenção (MF)"), ("coeficiente_utilizacao", "Coeficiente de Utilização (CU)"),
    ("tipo_ambiente", "Tipo Ambiente"), ("lux_requerido", "Lux Requerido"), ("qtd_luminarias", "Qtd Luminárias"),
    ("potencia_w", "Potência (W)"), ("fluxo_lumens", "Fluxo (Lumens)"), ("ip", "IP"),
    ("fluxo_total", "Fluxo Total"), ("eficiencia_lmxw", "Eficiência Luminosa (LmxW)"),
    ("temperatura_cor_k", "Temperatura Cor (K)"), ("tensao", "Tensão"), ("potencia_total_w", "Potência Total (W)"),
    ("densidade_w_m2", "Densidade (Wxm²)"), ("lux_calculado", "Lux Calculado"),
    ("diferenca_lux", "Diferença (lux Calc.-Lux Req.)"), ("atende", "Atende"),
]
_NCOLS = len(COLUNAS)


def gerar_excel_luminotecnico(projeto, dados: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Estudo Luminotécnico"
    linha = 1

    celula(ws, linha, 1, f"ESTUDO LUMINOTÉCNICO — {projeto.codigo_projeto or ''}", negrito=True,
           tamanho=14, alinhamento=ESQUERDA_LONGO, borda=False)
    linha += 2

    cabecalho_tabela(ws, linha, [rotulo for _, rotulo in COLUNAS])
    linha += 1
    for item in dados["linhas"]:
        for c, (chave, _) in enumerate(COLUNAS, start=1):
            v = item.get(chave)
            celula(ws, linha, c, v if v is not None else "—", alinhamento=ESQUERDA if chave in ("linha_succao", "ambiente") else CENTRO)
        linha += 1
    linha += 1

    barra(ws, linha, "RESUMO SOMATÓRIO — POR MODELO DE LUMINÁRIA", _NCOLS, cor_fundo=COR_TITULO)
    linha += 1
    cabecalho_tabela(ws, linha, ["Modelo Luminária", "Fabricante", "Qtd. Total"])
    linha += 1
    for r in dados["resumo_por_modelo"]:
        celula(ws, linha, 1, r["modelo_luminaria"], alinhamento=ESQUERDA)
        celula(ws, linha, 2, r.get("fabricante") or "—")
        celula(ws, linha, 3, r["qtd_total"])
        linha += 1
    linha += 1

    notas = dados.get("notas") or []
    if notas:
        celula(ws, linha, 1, "NOTAS TÉCNICAS", negrito=True, tamanho=12, alinhamento=ESQUERDA_LONGO, borda=False)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
        linha += 1
        for i, nota in enumerate(notas, start=1):
            celula(ws, linha, 1, f"{i}) {nota}", cor_fonte=COR_MUTED, tamanho=9,
                   alinhamento=ESQUERDA_LONGO, borda=False)
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=_NCOLS)
            linha += 1

    ws.freeze_panes = "A2"
    autosize(ws, _NCOLS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
