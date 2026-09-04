"""Importa a Normal Climatológica do Brasil 1991-2020 (INMET) para CondicaoClimatica.

Fontes (planilhas oficiais INMET, uma por grandeza):
  TMAX      -> tbs_pico_sazonal (maior valor mensal) e o mês em que ocorre
  TMEDUMID  -> tbu_pico_sazonal (no mesmo mês do pico de TMAX) e tbu_media_anual (coluna Ano)
  UR        -> ur_pico_sazonal (no mesmo mês do pico de TMAX) e ur_media_anual (coluna Ano)
  TMEDSECA  -> tbs_media_anual (coluna Ano — média verdadeira, não média dos máximos)
  TMAXABS   -> tbs_maxima_absoluta (recorde histórico da estação)
  TORV (ponto de orvalho) não é usado — a fórmula de infiltração usa TBS/TBU/UR diretamente.

Roda uma vez. Idempotente: apaga e recria os dados da tabela a cada execução.
"""
import sys
import os
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))
from backend.database import SessionLocal
from backend import models as m

PASTA = r"\\mycloudex2ultra\Téchne_Projetos\Biblioteca Projetos\Projetos\Planilhas Técnicas\Dados Climatologia"
MESES = list(range(3, 15))  # colunas 3..14 (0-indexed) = Janeiro..Dezembro nas abas de 16 colunas


def _num(v):
    if v is None or v == "-" or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _ler_simples(caminho, aba):
    """Lê as abas de 16 colunas (Código, Nome, UF, Jan..Dez, Ano). Devolve {codigo: {"cidade":..,
    "uf":.., "meses": [12 valores], "ano": valor}}."""
    wb = openpyxl.load_workbook(os.path.join(PASTA, caminho), data_only=True)
    ws = wb[aba]
    dados = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        codigo = row[0]
        if not codigo:
            continue
        meses = [_num(row[i]) for i in MESES]
        dados[codigo] = {"cidade": row[1], "uf": row[2], "meses": meses, "ano": _num(row[15])}
    return dados


def _ler_maxabs(caminho, aba):
    """TMAXABS tem estrutura diferente: por mês, um par (Ano, Valor); o último par é o
    recorde absoluto geral (Mês/Ano, Valor)."""
    wb = openpyxl.load_workbook(os.path.join(PASTA, caminho), data_only=True)
    ws = wb[aba]
    dados = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        codigo = row[0]
        if not codigo:
            continue
        dados[codigo] = {"cidade": row[1], "uf": row[2], "absoluta": _num(row[28])}
    return dados


def run():
    print("Lendo planilhas...")
    tmax = _ler_simples("Normal-Climatologica-TMAX.xlsx", "TMAX")
    tmedumid = _ler_simples("Normal-Climatologica-TMEDUMID.xlsx", "TMEDUMID")
    ur = _ler_simples("Normal-Climatologica-UR.xlsx", "UR")
    tmedseca = _ler_simples("Normal-Climatologica-TMEDSECA.xlsx", "Plan1")
    tmaxabs = _ler_maxabs("Normal-Climatologica-TMAXABS (1).xlsx", "TMAXABS")

    todos_codigos = set(tmax) | set(tmedumid) | set(ur) | set(tmedseca) | set(tmaxabs)
    print(f"{len(todos_codigos)} estações encontradas ao todo.")

    db = SessionLocal()
    db.query(m.CondicaoClimatica).delete()

    inseridas = 0
    for codigo in todos_codigos:
        t = tmax.get(codigo)
        if not t:
            continue  # sem TMAX não dá pra calcular o pico sazonal (critério padrão) — pula
        cidade, uf = t["cidade"], t["uf"]

        valores_mes = [(i, v) for i, v in enumerate(t["meses"]) if v is not None]
        if not valores_mes:
            continue
        idx_pico, tbs_pico = max(valores_mes, key=lambda x: x[1])

        tu = tmedumid.get(codigo)
        tbu_pico = tu["meses"][idx_pico] if tu and tu["meses"][idx_pico] is not None else None
        tbu_anual = tu["ano"] if tu else None

        u = ur.get(codigo)
        ur_pico = u["meses"][idx_pico] if u and u["meses"][idx_pico] is not None else None
        ur_anual = u["ano"] if u else None

        ts = tmedseca.get(codigo)
        tbs_anual = ts["ano"] if ts else None

        ta = tmaxabs.get(codigo)
        tbs_absoluta = ta["absoluta"] if ta else None

        db.add(m.CondicaoClimatica(
            cidade=cidade, uf=uf, codigo_estacao=codigo,
            tbs_pico_sazonal=tbs_pico, tbu_pico_sazonal=tbu_pico, ur_pico_sazonal=ur_pico,
            tbs_media_anual=tbs_anual, tbu_media_anual=tbu_anual, ur_media_anual=ur_anual,
            tbs_maxima_absoluta=tbs_absoluta,
        ))
        inseridas += 1

    db.commit()
    print(f"{inseridas} estações importadas em CondicaoClimatica.")
    db.close()


if __name__ == "__main__":
    run()
