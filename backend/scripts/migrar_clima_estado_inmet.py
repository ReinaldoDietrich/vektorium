# -*- coding: utf-8 -*-
"""Migração aditiva: cria estado_brasileiro (Tela 1 - Lista de Estados Brasileiros) e
dados_climatologicos_inmet (Tela 1 - Dados Climatológicos INMET 1990-2020), semeando das planilhas
"Tabela Lista Estados.xlsx" e "Tabela de Dados Climatológicos.xlsx". Adiciona projetos.estado_uf e
projetos.estacao_inmet_id. Substitui o seletor "Critério de Projeto" — fonte única de Temperatura/UR
passa a ser a Temp. Máxima Histórica (discussão real 2026-07-18).

Projetos já existentes ficam com estado_uf/estacao_inmet_id em branco (não dá pra inferir com
segurança qual estação nova corresponde à antiga — usuário re-seleciona Estado+Estação).

Idempotente. Rodar: python -m backend.scripts.migrar_clima_estado_inmet <caminho_db> [dir_planilhas]
"""
import sys
import sqlite3


def _colunas(cur, tabela):
    cur.execute(f"PRAGMA table_info({tabela})")
    return {row[1] for row in cur.fetchall()}


def _ler_estados(caminho_xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws = wb["Planilha1"]
    linhas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1]:
            continue
        linhas.append((row[0], row[1], row[2], row[3]))
    return linhas


def _ler_dados_climatologicos(caminho_xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(caminho_xlsx, data_only=True)
    ws = wb["TMAXABS"]
    linhas = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] is None:
            continue
        codigo, nome, uf, temp_max, ur_media = row[0], row[1], row[2], row[3], row[4]
        linhas.append((codigo, nome, uf, temp_max, ur_media))
    return linhas


def migrar(caminho_db, caminho_estados, caminho_dados_clima):
    con = sqlite3.connect(caminho_db)
    cur = con.cursor()

    cur.execute("""CREATE TABLE IF NOT EXISTS estado_brasileiro (
        id INTEGER PRIMARY KEY,
        estado VARCHAR NOT NULL,
        sigla VARCHAR NOT NULL UNIQUE,
        capital VARCHAR,
        regiao VARCHAR
    )""")
    cur.execute("SELECT COUNT(*) FROM estado_brasileiro")
    if cur.fetchone()[0] == 0:
        estados = _ler_estados(caminho_estados)
        cur.executemany("INSERT INTO estado_brasileiro (estado, sigla, capital, regiao) VALUES (?, ?, ?, ?)", estados)
        print(f"estado_brasileiro: {len(estados)} linha(s) semeada(s).")
    else:
        print("estado_brasileiro já tinha dados — não sobrescrita.")

    cur.execute("""CREATE TABLE IF NOT EXISTS dados_climatologicos_inmet (
        id INTEGER PRIMARY KEY,
        codigo INTEGER NOT NULL,
        nome_estacao VARCHAR NOT NULL,
        uf VARCHAR NOT NULL,
        temp_maxima_historica FLOAT,
        ur_media_historica FLOAT
    )""")
    cur.execute("SELECT COUNT(*) FROM dados_climatologicos_inmet")
    if cur.fetchone()[0] == 0:
        dados = _ler_dados_climatologicos(caminho_dados_clima)
        cur.executemany(
            "INSERT INTO dados_climatologicos_inmet (codigo, nome_estacao, uf, temp_maxima_historica, ur_media_historica) "
            "VALUES (?, ?, ?, ?, ?)", dados)
        print(f"dados_climatologicos_inmet: {len(dados)} linha(s) semeada(s).")
    else:
        print("dados_climatologicos_inmet já tinha dados — não sobrescrita.")

    colunas_projetos = _colunas(cur, "projetos")
    if "estado_uf" not in colunas_projetos:
        cur.execute("ALTER TABLE projetos ADD COLUMN estado_uf VARCHAR")
        print("projetos.estado_uf criada.")
    else:
        print("projetos.estado_uf já existia.")
    if "estacao_inmet_id" not in colunas_projetos:
        cur.execute("ALTER TABLE projetos ADD COLUMN estacao_inmet_id INTEGER")
        print("projetos.estacao_inmet_id criada.")
    else:
        print("projetos.estacao_inmet_id já existia.")

    con.commit()
    con.close()
    print("Migração concluída.")


if __name__ == "__main__":
    caminho = sys.argv[1] if len(sys.argv) > 1 else r"B:\Documentos Programas\App Carga Térmica\database\carga_termica.db"
    diretorio = sys.argv[2] if len(sys.argv) > 2 else r"B:\Documentos Programas\App Carga Térmica\Documentos de Criação"
    migrar(caminho, f"{diretorio}\\Tabela Lista Estados.xlsx", f"{diretorio}\\Tabela de Dados Climatológicos.xlsx")
